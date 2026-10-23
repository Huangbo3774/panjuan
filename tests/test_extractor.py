"""题型识别单测：正例 + 反例。

反例尤其重要：PRD 的题号正则 `^\\d{1,3}[.、．)]` 会把 "4. 4×7 读作" 里的
"4×7" 当成新题号，也会把 "1." 之外的乘号误判。这里把这类坑钉住。
"""
from __future__ import annotations

import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from app.extractor.rules import extract, flatten_questions            # noqa: E402
from app.extractor.rules import split_subqs                          # noqa: E402
from app.extractor.rules import _parse_question_no as _parse          # noqa: E402
from app.extractor.lines import split_question_ann, spans_for_questions  # noqa: E402

ok_n = bad_n = 0


def t(name, fn):
    global ok_n, bad_n
    try:
        fn(); ok_n += 1; print("  PASS " + name)
    except AssertionError as e:
        bad_n += 1; print(f"  FAIL {name} → {e}")
    except Exception as e:
        bad_n += 1; print(f"  ERROR {name} → {type(e).__name__}: {e}")


def qs(res):
    return flatten_questions(res)


print("\n[基本切分]")
def basic():
    lines = [
        "一、选择题（每题 3 分，共 15 分）",
        "1. 4×7 读作（  ）",
        "A. 4乘7    B. 4乘以7    C. 7乘4    D. 四七二十八",
        "2. 一个因数是 6，另一个因数是 8，积是（  ）",
        "A. 14      B. 48        C. 24      D. 16",
        "二、填空题（每空 2 分，共 20 分）",
        "3. 一年有＿＿个月，其中大月有＿＿个。",
        "三、解答题（共 25 分）",
        "4. 果园里有 6 行桃树，每行 8 棵，一共有多少棵？",
    ]
    r = extract(lines, "测试卷")
    q = qs(r)
    assert [x["q_no"] for x in q] == [1, 2, 3, 4], f"题号错: {[x['q_no'] for x in q]}"
    assert q[0]["qtype"] == "choice", f"第1题应 choice，实际 {q[0]['qtype']}"
    assert len(q[0]["options"]) == 4, f"选项数 {len(q[0]['options'])}: {q[0]['options']}"
    assert q[2]["qtype"] == "fill", f"第3题应 fill，实际 {q[2]['qtype']}"
    assert q[2]["blanks"] == 2, f"应有 2 个空，实际 {q[2]['blanks']}"
    assert q[3]["qtype"] == "essay", f"第4题应 essay，实际 {q[3]['qtype']}"
    assert len(r["sections"]) == 3, f"大题数 {len(r['sections'])}"
t("选择/填空/解答三题型 + 3 个大题", basic)


def multiply_not_question():
    """反例：题干里的乘号不能被当成题号（PRD 原正则的坑）。"""
    lines = ["一、选择题", "1. 4×7 读作（  ）", "A. x    B. y    C. z    D. w"]
    q = qs(extract(lines))
    assert len(q) == 1, f"应只切出 1 题，实际 {len(q)}: {[x['text'] for x in q]}"
t("题干中的乘号不被误判为题号（反例）", multiply_not_question)


def decimal_not_question():
    """反例：句点后不要求空格后，小数不能被当成题号。"""
    assert _parse("3.5 与 4.2 哪个大") == (None, "3.5 与 4.2 哪个大")
    assert _parse("1.一共有（ ）个") == (1, "一共有（ ）个")
    assert _parse("长 3.5 厘米") == (None, "长 3.5 厘米")
t("小数不被当题号，句点无空格仍可切（反例）", decimal_not_question)


def inline_options():
    lines = ["一、选择题", "3. 五八（  ）四九（  ）六七（  ）",
             "A. 四十 三六 四十二     B. 四十 三十六 四十二",
             "C. 四十五 三十六 四十二  D. 四十 三十六 四十八"]
    q = qs(extract(lines))
    assert len(q) == 1, f"应 1 题，实际 {len(q)}"
    assert len(q[0]["options"]) >= 3, f"选项太少: {q[0]['options']}"
    assert q[0]["qtype"] == "choice"
t("一行内多个选项也能拆开", inline_options)


def judge_type():
    lines = ["二、判断题（对的打√，错的打×）",
             "1. 0 乘任何数都得 0。",
             "2. 5×0 和 0+5 结果相同。"]
    q = qs(extract(lines))
    assert all(x["qtype"] == "judge" for x in q), f"应全 judge: {[x['qtype'] for x in q]}"
t("判断题识别（对的打√）", judge_type)


def multiplication_is_not_a_judge_signal():
    """乘号不是判断信号：题干里只要有 × 就判 judge，做题页只给对/错，孩子没处写答案。

    真实卷（第三单元练习题）实测 8 道乘法题因此被判成判断题；用户原话"这题都怎么做啊"。
    判断题的信号必须是"对的打/判断/画√"这类指令，不是光秃秃一个符号。
    """
    lines = ["一、填一填。",
             "1. 2×3 可以表示（  ）个（  ）相加。",
             "2. 6×2=（  ），5×3=（  ）。"]
    q = qs(extract(lines))
    assert len(q) == 2, f"应切出 2 题：{[x['text'][:20] for x in q]}"
    assert all(x["qtype"] != "judge" for x in q), \
        f"乘法题被判成判断题：{[(x['text'][:16], x['qtype']) for x in q]}"
    assert q[0]["qtype"] == "fill" and q[1]["qtype"] == "fill", \
        f"应是有空位的填空题：{[x['qtype'] for x in q]}"
t("乘法题不许被判成判断题", multiplication_is_not_a_judge_signal)


def picture_choice_is_flagged_not_turned_into_judge():
    """选项是图片的选择题：OCR 读不到 A/B/C/D，不许判成判断题，且要挂待校对。

    真实卷第 3 题（自然科普·树叶）：题干"不能用2×4表示的是（ ）(填序号)"，
    四片叶子是图片选项。修掉乘号误判后它应是可作答的 fill，并提醒家长补选项。
    """
    lines = ["第2天 儿童乐园",
             "3. 下面的树叶中，不能用2×4表示的是（  ）。(填序号)",
             "4. 第二题的题干文字。"]
    q = qs(extract(lines))
    assert len(q) == 2, f"应切出 2 题：{[x['text'][:18] for x in q]}"
    assert q[0]["qtype"] != "judge", f"又是判断题：{q[0]['qtype']}"
    assert q[0]["qtype"] == "fill", f"应保留成可作答的 fill：{q[0]['qtype']}"
    assert q[0]["needs_review"], "题干明写填序号却没选项，必须挂待校对"
    assert "选择题" in str(q[0].get("review_note") or ""), \
        f"待校对理由没说清：{q[0].get('review_note')}"
t("图片选项的选择题：可作答 + 待校对", picture_choice_is_flagged_not_turned_into_judge)


def fill_without_blanks_becomes_essay():
    """图内文字清理会把只印在图里的空位删掉（P5 Q12/Q16）：题型还是 fill，
    但空位 0 个。做题页按 fill 渲染就一个输入框都不给 —— 学生没法作答。
    末尾重算：fill 且无空位 → essay（有作答框）；真有空位的 fill 不许动。
    """
    from app.extractor.rules import recount_blanks
    q = [
        {"qtype": "fill", "text": "观察下图，分别用不同的数表示图中阴影部分占整图的多少",
         "images": [{"name": "x.png"}]},
        {"qtype": "fill", "text": "这扇窗户的面积是（    ）平方米。"},
        {"qtype": "choice", "text": "下面正确的是（  ）", "options": ["A.1", "B.2"]},
        {"qtype": "essay", "text": "说说你的理由。"},
    ]
    assert recount_blanks(q) == 1
    assert q[0]["qtype"] == "essay" and q[0]["blanks"] == 0, q[0]
    assert q[1]["qtype"] == "fill", "有空位的填空题被误改：" + q[1]["qtype"]
    assert q[2]["qtype"] == "choice" and q[3]["qtype"] == "essay"
t("空位清光后的 fill 退成 essay（含反例）", fill_without_blanks_becomes_essay)


def uncertain_types_fall_back_to_essay_answer_box():
    """题型判不清/复杂到装不下 → 退成问答题，保证有一个自由作答框。

    用户口径：复杂题判断不清题型时，不硬塞成单选或一个判断空。只动四类：
    unknown、choice 选项不足两个、选择/判断带多个小问、填空小问多于空位。
    真选择（≥2 选项）、真判断、空位齐全的填空一律不动（反例）。
    """
    from app.extractor.rules import fallback_uncertain_types
    q = [
        {"qtype": "unknown", "text": "松树种了多少棵？", "options": [], "blanks": 0},
        {"qtype": "choice", "text": "商是几位数？", "options": ["①4"],
         "blanks": 1},
        {"qtype": "choice", "text": "下面哪幅图正确？", "options": ["A", "B", "C"],
         "subqs": [{"no": 1, "text": "圈一圈"}, {"no": 2, "text": "写一写"}]},
        {"qtype": "judge", "text": "判断各题。", "options": [],
         "subqs": [{"no": 1, "text": "甲"}, {"no": 2, "text": "乙"}]},
        {"qtype": "fill", "text": "填一填。", "blanks": 1,
         "subqs": [{"no": 1, "text": "甲"}, {"no": 2, "text": "乙"}]},
        # 反例：真选择、真判断、空位齐全的填空
        {"qtype": "choice", "text": "选一选。", "options": ["A.1", "B.2", "C.3"]},
        {"qtype": "judge", "text": "对的画√，错的画×。", "options": [], "blanks": 1},
        {"qtype": "fill", "text": "窗户面积是（    ）平方米。", "blanks": 1},
        {"qtype": "essay", "text": "说说你的理由。"},
    ]
    assert fallback_uncertain_types(q) == 5
    assert all(x["qtype"] == "essay" for x in q[:5]), \
        f"判不清题型的题没退成问答：{[x['qtype'] for x in q[:5]]}"
    assert all(x["needs_review"] for x in q[:5]), "退成问答后没挂待校对"
    assert all("问答题" in str(x.get("review_note") or "") for x in q[:5]), \
        "没告诉家长为什么改成问答"
    assert [x["qtype"] for x in q[5:]] == ["choice", "judge", "fill", "essay"], \
        f"正常题型被误改：{[x['qtype'] for x in q[5:]]}"


t("判不清题型退成问答框（含四个反例）", uncertain_types_fall_back_to_essay_answer_box)


def real_judge_section_still_works():
    """反例保险：修掉乘号后，真正的判断题（大题名给的信号）不能跟着丢。"""
    lines = ["三、判断。（对的画√，错的画×）",
             "1. 所有直角都一样大。",
             "2. 长方形的四条边都相等。"]
    q = qs(extract(lines))
    assert all(x["qtype"] == "judge" for x in q), f"判断题丢了：{[x['qtype'] for x in q]}"
t("真判断题仍然认（反例保险）", real_judge_section_still_works)


def judge_word_inside_essay_is_not_judge():
    """正文里出现"判断"两个字，不等于这道题是判断题。

    P5 Q24「请做出判断并说明理由」、Q25「请你判断：这三个图形的周长相等吗？」
    都是多问解答题。旧口径见"判断"就判 judge，导出时在题末凭空印
    "判断：（    ）"——一道不存在的判断题，孩子不知道答什么，老师也没法批。
    """
    lines = ["五、解决问题（共5道小题）",
             "24. 张老师记录了六（1）班和六（2）班男生的成绩。"
             "请按成绩段整理数据，补全下表，并做出判断、说明理由。",
             "25. 同学们设计出了三个新图形，请你判断：这三个图形的周长相等吗？"]
    q = qs(extract(lines))
    assert len(q) == 2, f"应切出 2 题：{[x['text'][:20] for x in q]}"
    assert all(x["qtype"] != "judge" for x in q), \
        f"正文里的'判断'把解答题判成了判断题：{[(x['text'][:18], x['qtype']) for x in q]}"
    assert all(x["qtype"] == "essay" for x in q), \
        f"应是无选项的解答题：{[x['qtype'] for x in q]}"
t("正文含「判断」的解答题不许判成判断题", judge_word_inside_essay_is_not_judge)


def header_not_question():
    """反例：表头「班级/姓名/满分」不该变成题目。"""
    lines = ["第三单元达标练习", "班级＿＿ 姓名＿＿ 得分＿＿", "满分：100 分  时间：40 分钟",
             "一、选择题", "1. 例子（  ）", "A. a  B. b  C. c  D. d"]
    r = extract(lines)
    q = qs(r)
    assert len(q) == 1, f"应只有 1 题，实际 {len(q)}: {[x['text'] for x in q]}"
    assert any("班级" in x for x in r["skipped_header_lines"]), "表头应被记为跳过行"
t("表头不当题目（反例）", header_not_question)


def no_loss():
    """核心不变式：一行内容都不能丢。宁可切错也不能漏。"""
    lines = ["一、选择题", "1. 题目甲（  ）", "A. a  B. b  C. c  D. d",
             "2. 题目乙。", "3. 题目丙。", "二、填空题", "4. 空是＿＿。",
             "5. 长题干题目，需要写出计算过程和答语。"]
    r = extract(lines)
    joined = " ".join(x["text"] + " " + " ".join(x["options"]) for x in qs(r))
    for key in ("题目甲", "题目乙", "题目丙", "空是", "计算过程"):
        assert key in joined, f"内容丢失：{key}"
t("内容零丢失不变式", no_loss)


def empty_paren_is_blank():
    """PRD 只写「全角括号长度≥2」；二年级的（　）常是 0~2 个空格，也要认成空位。"""
    lines = ["一、填空题", "1. 光沿（）传播。", "2. 声音速度（  ）米每秒。"]
    q = qs(extract(lines))
    assert q[0]["qtype"] == "fill", f"空括号应判 fill，实际 {q[0]['qtype']}"
    assert q[1]["qtype"] == "fill"
t("空括号也算填空位（补 PRD 缺口）", empty_paren_is_blank)


def never_empty():
    """兜底：完全没有题号的乱文本，也必须返回至少一道题，不能返回空。
    （加了"隐式大题"之后，这类行走正常切题路径而非兜底路径，
      所以断言改为"内容不丢 + 有题产出"，不再要求置信度极低。）"""
    r = extract(["这是一段没有任何题号的文字内容"])
    q = qs(r)
    assert len(q) >= 1, "兜底失败，返回了 0 题"
    assert "文字内容" in q[0]["text"], f"内容丢失: {q[0]['text']}"
    # 必须走【隐式大题】这条正常路径，不许退到 fallback：
    # fallback 把每行都压成 confidence 0.2 的"整页当一堆文字"，是最后一道保险而不是常态。
    # 这条断言也让「没有大题标题就开一节」变成一把被量过的尺（变异闸门验证过它会红）。
    assert not any(x.get("fallback") for x in q), f"退到了 fallback 路径: {[x.get('confidence') for x in q]}"
    # 真正走到兜底分支时（全空白输入）才要求极低置信度
    r2 = extract(["   ", "", "  "])
    assert qs(r2) == [], "全空白输入不该造出题"
t("无题号文本不丢内容（反例）", never_empty)


def confidence_gate():
    lines = ["一、选择题", "1. 例子（  ）", "A. a  B. b  C. c  D. d", "2. 短"]
    q = qs(extract(lines))
    assert q[0]["needs_review"] is False, "清晰选择题不该标待校对"
    assert any(x["needs_review"] for x in q), "低置信题应标待校对"
t("置信度 <0.7 标待校对", confidence_gate)


def circled_numbers():
    """连线/排序题里 ①②③ 是条目 → 当题号。"""
    lines = ["四、排序题", "①最后用帽子盖住鸡蛋顶上的小洞。", "②然后用彩纸做一顶尖尖的小帽子。",
             "③先在鸡蛋顶部打一个小洞。"]
    q = qs(extract(lines))
    assert len(q) == 3, f"排序题里圆圈序号应切成 3 项，实际 {len(q)}"
    assert [x["q_no"] for x in q] == [1, 2, 3], f"{[x['q_no'] for x in q]}"
t("排序题里圆圈数字作题号", circled_numbers)


def circled_are_options_in_choice():
    """反例：选择题里 ①②③ 是选项，不能当题号（扫描卷实测：会被切成 4 道假题）。"""
    lines = ["一、选择题", "1. 小宇在家的（ ）方向。", "①东 ②南 ③西 ④北"]
    q = qs(extract(lines))
    assert len(q) == 1, f"应只 1 题，实际 {len(q)}: {[x['text'] for x in q]}"
    assert len(q[0]["options"]) == 4, f"圆圈应成为 4 个选项，实际 {q[0]['options']}"
    assert q[0]["qtype"] == "choice", f"应 choice，实际 {q[0]['qtype']}"
t("选择题里圆圈数字作选项（反例）", circled_are_options_in_choice)


def circled_options_written_inline():
    """正例：选项和题干挤在同一段（真实卷 第三单元综合练习 二、选一选 第 1/2 题）。

    OCR 把它们连成 "…错误的是（ ①30表示… ②5表示… ③6表示…"，从前只认行首圆圈，
    于是整道题判成 essay 并报"题型可疑"；这里必须拆成 choice + 3 个选项。
    """
    lines = ["二、选一选。",
             "1.根据下图可列出乘法算式6×5=30,下列说法中,错误的是（",
             "①30表示一共有30支笔",
             "②5表示每个笔筒中有5支笔",
             "③6表示每个笔筒中有6支笔"]
    q = qs(extract(lines))
    assert len(q) == 1, f"应只 1 题，实际 {len(q)}: {[x['text'] for x in q]}"
    assert q[0]["qtype"] == "choice", f"应 choice，实际 {q[0]['qtype']}"
    assert len(q[0]["options"]) == 3, f"应拆出 3 个选项，实际 {q[0]['options']}"
    assert "①" not in q[0]["text"] and "③" not in q[0]["text"], \
        f"选项还留在题干里：{q[0]['text']}"
t("题干内联圆圈选项要拆出来", circled_options_written_inline)


def circled_stay_items_in_matching_section():
    """反例：连线/排序题里的 ①②③ 是条目，内联拆分不能把它们当选项吞掉。"""
    lines = ["四、排序题",
             "①最后用帽子盖住鸡蛋顶上的小洞。",
             "②然后用彩纸做一顶尖尖的小帽子。",
             "③先在鸡蛋顶部打一个小洞。"]
    q = qs(extract(lines))
    assert len(q) == 3, f"应 3 项，实际 {len(q)}"
    assert all(not x["options"] for x in q), f"条目被当选项：{[x['options'] for x in q]}"
t("连线/排序的圆圈条目不拆成选项（反例）", circled_stay_items_in_matching_section)


def map_label_instruction_is_not_choice():
    """反例：「把…序号填在上图的括号里」是地图标注，不是选择题，不许挂待校对。"""
    lines = ["五、小玉绘制了自己家附近的平面图。",
             "1.根据描述，把建筑物的序号填在上图的括号里。(12分)",
             "我家住在祥和小区，南面是春风小区。"]
    q = qs(extract(lines))
    assert q[0]["qtype"] != "choice", f"地图标注被改成了 choice：{q[0]['qtype']}"
    assert not q[0].get("review_note"), f"地图标注被标成题型可疑：{q[0].get('review_note')}"
t("「序号填在图上」不是选择题信号（反例）", map_label_instruction_is_not_choice)


def answer_key_not_questions():
    """反例：答案区不能被切成题目（实测英文样本 "参考答案：1.B 2.A" 被切成 7 道假题）。"""
    lines = ["一、选择题", "1. 例子（ ）", "A. a  B. b  C. c  D. d",
             "2. 例子二（ ）", "A. e  B. f  C. g  D. h",
             "参考答案：1.B 2.A"]
    r = extract(lines)
    q = qs(r)
    assert len(q) == 2, f"应只 2 题，实际 {len(q)}: {[(x['q_no'],x['text']) for x in q]}"
    assert r["answer_key_lines"], "答案区应被单独收集"
    assert any("1.B" in x or "B" in x for x in r["answer_key_lines"])
t("答案区不当题目（反例）", answer_key_not_questions)


def lone_letter_not_question():
    """反例：单独成行的 "B" 是答案，不是第 2 题。"""
    lines = ["一、选择题", "1. 例子（ ）", "A. a  B. b  C. c  D. d", "B"]
    q = qs(extract(lines))
    assert len(q) == 1, f"应只 1 题，实际 {len(q)}"
t("孤立字母不当题（反例）", lone_letter_not_question)


print("\n[分段·水印·栏目（真实扫描卷暴露的缺口）]")


def _dup(res):
    """重号只在大题内部算：中文卷每大题都从 1 重新编号。"""
    n = 0
    for sec in res["sections"]:
        sn = [q["q_no"] for q in sec["questions"] if q["q_no"] is not None]
        n += len(sn) - len(set(sn))
    return n


def section_by_ordinal_not_keyword():
    """正例：大题名既不带动词的关键词（「三、公园畅游指南」「四、应急疏散路线指引」
    是真实卷的原话），只能靠【序号递增】认出分段。关键词表永远补不全。"""
    lines = ["一、填一填。", "1.甲（  ）", "2.乙（  ）",
             "三、公园畅游指南。（共21分）", "1.丙（  ）", "①东  ②西  ③北",
             "四、应急疏散路线指引。（共18分）", "2.丁（  ）"]
    r = extract(lines)
    names = [x["name"] for x in r["sections"]]
    assert len(r["sections"]) == 3, f"应 3 个大题，实际 {len(r['sections'])}: {names}"
    assert _dup(r) == 0, f"重号 {_dup(r)}: {names}"
    assert any("公园畅游" in n for n in names), f"无关键词大题没被认出: {names}"
t("无关键词大题靠序号递增切开", section_by_ordinal_not_keyword)


def increasing_numbers_do_not_split():
    """反例：同一小题号序列里编号递增，不能被"倒退分段"误切。"""
    lines = ["一、选择题", "1.题A（  ）", "A. a  B. b", "2.题B（  ）", "A. a  B. b", "3.题C（  ）"]
    r = extract(lines)
    assert len(r["sections"]) == 1, f"应 1 个大题，实际 {len(r['sections'])}"
    assert len(qs(r)) == 3, f"应 3 题，实际 {len(qs(r))}"
t("编号递增不误分段（反例）", increasing_numbers_do_not_split)


def restart_splits_without_header():
    """正例：编号从 3 倒回 1 且中间没有标题行 → 新的编号序列（教辅书每天重新编号）。"""
    lines = ["第1天", "1.甲", "2.乙", "第三步提思维", "3.丙", "1.丁", "2.戊"]
    r = extract(lines)
    assert _dup(r) == 0, f"重号 {_dup(r)}，分段 {[x['name'] for x in r['sections']]}"
    assert len(r["sections"]) == 2, f"应 2 段，实际 {len(r['sections'])}"
t("编号倒退自动分段", restart_splits_without_header)


def noisy_day_header_is_section():
    """回归：标题后面拖了水印长尾（剩余 27 字）也必须认出分段。
    曾用「剩余长度 ≤24 才算分段」，真实行「第2天 画校园路线图[课时2]   拍照秒批」漏认，
    导致整卷重号 20。"""
    lines = ["第1天 认识方向[课时1]            拍照秒批", "1.甲（  ）",
             "第2天   画校园路线图[课时2]            拍照秒批          视频讲解", "1.乙（  ）"]
    r = extract(lines)
    assert len(r["sections"]) == 2, f"应 2 段，实际 {len(r['sections'])}: {[x['name'] for x in r['sections']]}"
    assert _dup(r) == 0, f"重号 {_dup(r)}"
    assert "拍照秒批" not in r["sections"][0]["name"], f"标题没洗净水印: {r['sections'][0]['name']}"
t("带水印的日/课时标题仍认分段", noisy_day_header_is_section)


def watermark_lines_are_dropped():
    """正例：扫描 App 水印、页码、跑马页眉进 dropped_noise，不进题目。"""
    lines = ["一、填空题", "1.甲（  ）", "夸克扫描王", "极速扫描，就是高效", "·38·",
             "小学数学二年级上册·BS版", "2.乙（  ）"]
    r = extract(lines)
    q = qs(r)
    assert len(q) == 2, f"应 2 题，实际 {len(q)}: {[x['text'] for x in q]}"
    assert not any("扫描" in x["text"] for x in q), "水印混进题干"
    assert len(r["dropped_noise"]) == 4, f"噪声计数 {len(r['dropped_noise'])}: {r['dropped_noise']}"
t("水印/页码/页眉不进题目", watermark_lines_are_dropped)


def bare_number_not_pageno():
    """反例：「2.」是题号后内容全在图里的残缺题，绝不能当页码丢掉。"""
    lines = ["一、填空题", "1.甲（  ）", "2.", "3.丙（  ）"]
    r = extract(lines)
    got = [q["q_no"] for q in qs(r)]
    assert 2 in got, f"题号 2 被当页码丢了: {got}"
    assert "2." not in r["dropped_noise"], f"「2.」进了噪声: {r['dropped_noise']}"
t("裸题号「2.」不当页码（反例）", bare_number_not_pageno)


def stage_marker_not_question():
    """栏目行（第一步过考点 / 知识演练场）既不是题目也不是分段——编号会跨栏目延续。"""
    lines = ["第1天", "第一步过考点", "1.甲", "第二步练变式", "2.乙", "第三步提思维", "3.丙"]
    r = extract(lines)
    assert len(qs(r)) == 3, f"栏目行变成了题目: {[x['text'] for x in qs(r)]}"
    assert len(r["sections"]) == 1, f"栏目被误当分段: {[x['name'] for x in r['sections']]}"
    assert len(r["stage_markers"]) == 3, f"栏目计数 {len(r['stage_markers'])}"
t("教辅栏目行不当题/不当分段（反例）", stage_marker_not_question)


def english_section_detected():
    """正例：英文卷 Section A/B/C 要认作分段（曾用错正则分组，永远认不出）。"""
    lines = ["Section A  Multiple Choice", "1. one（  ）", "A. a  B. b",
             "Section B  Fill in the blank", "2. two __"]
    r = extract(lines)
    names = [x["name"] for x in r["sections"]]
    assert len(r["sections"]) == 2, f"应 2 段，实际 {names}"
    assert any("Multiple Choice" in n for n in names), names
t("英文 Section 识别为分段", english_section_detected)


def ordinal_guard_keeps_stem():
    """反例：以「二、」开头但自带答题空位/很长的行是题干，不能被序号判据吞成大题标题。"""
    lines = ["一、填空题", "1.甲（  ）", "二、把下面的数填在（  ）里：24、6、30、48"]
    r = extract(lines)
    assert len(r["sections"]) == 1, f"题干被当大题: {[x['name'] for x in r['sections']]}"
t("长题/带空位的行不当大题（反例）", ordinal_guard_keeps_stem)


def long_title_accepted():
    """回归：真实大题名可以带一长句引入（31 字）。曾用 24/30 字上限，
    把「六、随着无人驾驶技术的发展，无人配送车已投入使用。(共 12 分)」吞掉，
    该大题的题被并进上一节，报成"·续"+假重号。"""
    lines = ["五、小玉绘制了自己家附近的平面图。(共 18 分)", "1.甲（  ）",
             "六、随着无人驾驶技术的发展，无人配送车已投入使用。(共 12 分)", "1.乙（  ）"]
    r = extract(lines)
    names = [x["name"] for x in r["sections"]]
    assert len(r["sections"]) == 2, f"应 2 个大题，实际 {names}"
    assert _dup(r) == 0, f"重号 {_dup(r)}: {names}"
    assert not any("·续" in n for n in names), f"标题被吞: {names}"
t("长引入句的大题名仍认分段（回归）", long_title_accepted)


def bonus_section_keeps_its_own_question():
    """「附加题。(共10分)」必须另起一段，不能并进上一题。

    真实卷（综合实践单元(1) p4）：不认这段标题，附加题的两行填空会被
    上一题的区间吞掉，配图一路裁到页脚，把附加题整段做成位图。
    """
    lines = ["六、解决问题。", "1.甲（  ）", "2.乙（  ）",
             "从标“★”的快递点出发，需要经过每个小区。",
             "附加题。(共10分)",
             "山的南面与水的北面光照更充足，故称“阳”。",
             "“贵阳”因位于贵山的（  ）面而得名。"]
    r = extract(lines)
    names = [x["name"] for x in r["sections"]]
    assert names[-1] == "附加题。(共10分)", f"附加题没成段: {names}"
    prev = r["sections"][-2]["questions"][-1]
    assert "附加题" not in prev["text"], f"标题并进上一题: {prev['text']!r}"
    assert "贵阳" not in prev["text"], f"附加题正文并进上一题: {prev['text']!r}"
    bonus = r["sections"][-1]["questions"][0]
    assert bonus["q_no"] is None, f"附加题不该有题号: {bonus['q_no']}"
    assert bonus["auto_no"] == 1, f"附加题在大题内应排第 1: {bonus.get('auto_no')}"
    assert "贵阳" in bonus["text"], f"附加题正文丢了: {bonus['text']!r}"
t("附加题单独成段（含反例）", bonus_section_keeps_its_own_question)


def bonus_like_stem_not_split():
    """反例：题号后面/正文中间提到「附加题」的句子不是分段标题。"""
    lines = ["一、填空。", "1.完成附加题。(共10分)（  ）", "2.乙（  ）",
             "附加题部分共有三道小题，请认真作答。"]
    r = extract(lines)
    names = [x["name"] for x in r["sections"]]
    assert names == ["一、填空。"], f"题干里的附加题被当分段: {names}"
    assert len(r["sections"][0]["questions"]) == 2, "题目被切坏"
    assert "附加题部分" in r["sections"][0]["questions"][1]["text"], "续行正文被吞"
t("题干里的附加题不误分段（反例）", bonus_like_stem_not_split)


def long_essay_confidence():
    """正例：完整长问句（够长 + 句末标点）置信度要过 0.7，别把整卷标成待校对。
    反例：短残句（OCR 断行/图形碎片）必须仍然标 needs_review。"""
    lines = ["四、解决问题。",
             "1.果园里有 6 行桃树，每行 8 棵，一共有多少棵？请写出算式并作答。",
             "2.东", "3.图书馆       西            操场"]
    r = extract(lines)
    q = qs(r)
    assert q[0]["qtype"] == "essay", q[0]["qtype"]
    assert not q[0]["needs_review"], f"完整长题干被标待校对: {q[0]['confidence']}"
    assert all(x["needs_review"] for x in q[1:]), f"短残句没标待校对: {[(x['text'],x['confidence']) for x in q[1:]]}"
t("长问答置信达标/短残句仍待校对", long_essay_confidence)


def paper_title_promoted():
    """正例：卷名行（没题号、没空位、短）不该变成第 1 道题。"""
    lines = ["第三单元 过关练习", "班级＿＿ 姓名＿＿ 满分 100",
             "一、选择题", "1. 4×7 读作（  ）", "A. 四七二十八  B. 4乘7"]
    r = extract(lines)
    q = qs(r)
    assert len(q) == 1, f"卷名/表头变成了题目: {[(x['q_no'], x['text'][:20]) for x in q]}"
    assert r["title"] == "第三单元 过关练习", r["title"]
    assert any("班级" in x for x in r["skipped_header_lines"]), r["skipped_header_lines"]
t("卷名提升为标题、不当题目", paper_title_promoted)


def short_first_line_with_blank_stays():
    """反例：带空位的首行是题目（"填一填：（ ）"），不能被当成卷名吃掉。"""
    lines = ["填一填：（  ）", "1. 甲（  ）"]
    r = extract(lines)
    assert len(qs(r)) == 2, [x["text"] for x in qs(r)]
    assert not r["title"], f"带空位的行被当卷名: {r['title']}"
t("带空位的首行仍是题目（反例）", short_first_line_with_blank_stays)


def margin_strip_never_enters_body_text():
    """装订栏（竖排"学校/班级/姓名/学号"、"弥/封/线"）必须在【合并成行之前】被摘掉。

    真实卷：综合实践单元.pdf 841×653 横版两张 A4，左右各一条装订栏。
    这些竖排单字和正文同一水平线，一旦进了 group_lines 就变成
    「1. 学号： 线 小朋友们，准备好…」「…后面是（ 封」——用户说的"排版变形"根子在这。
    """
    from app.extractor.lines import drop_margin_items
    items = [
        {"text": "学号：", "x": 10, "y": 157, "w": 13, "h": 13},
        {"text": "线", "x": 27, "y": 145, "w": 9, "h": 9},
        {"text": "弥", "x": 26, "y": 426, "w": 9, "h": 11},
        {"text": "密封线内不要答题", "x": 800, "y": 300, "w": 30, "h": 90},
        {"text": "综合实践", "x": 800, "y": 140, "w": 34, "h": 92},      # 右侧栏色块
        {"text": "1. 小朋友们在操场上辨认方向，请帮他们补全记录单。",
         "x": 64, "y": 315, "w": 287, "h": 14},
    ]
    keep, dropped = drop_margin_items(items, 841.0)
    assert [d["text"] for d in dropped] == ["学号：", "线", "弥", "密封线内不要答题", "综合实践"], \
        f"装订栏没被清干净/清错：{[(d['x'], d['text']) for d in dropped]}"
    assert len(keep) == 1 and keep[0]["text"].startswith("1. 小朋友"), keep


def margin_filter_must_not_eat_body_line_ends():
    """反例：贴边但属于正文的行尾碎片不能丢。

    横版卷左半正文就从 x≈45 起笔。第一版用"框左边 ≤45"判贴边，
    于是把题干里短到两三个字的行尾（"）。"）也当成装订栏吃掉，
    用户看到的是"…一端指(  ）"读不通——比混进装订栏更糟。
    现在要求【整个框】都在页边带里。
    """
    from app.extractor.lines import drop_margin_items
    body = [{"text": "）。", "x": 40, "y": 240, "w": 26, "h": 12},
            {"text": "）；", "x": 30, "y": 260, "w": 18, "h": 12}]
    keep, dropped = drop_margin_items(body, 841.0)
    assert not dropped, f"正文行尾被当装订栏吃了：{[(d['x'], d['text']) for d in dropped]}"
    assert len(keep) == 2
    strip = [{"text": "学校：", "x": 8, "y": 470, "w": 22, "h": 12}]
    _k, dr = drop_margin_items(strip, 841.0)
    assert len(dr) == 1, "真装订栏反而没抓到"


def sealing_column_on_photo_edge_is_dropped():
    """照片右边缘的密封线整列必须摘掉，不能只靠 45pt 页边。

    真实卷：第三单元达标练习题B.jpg（1250×881px ⇒ 900×634pt），
    OCR 把「密封线内请勿答题」读成十个单字框，散落在 x=838–860pt，
    其中"请/勿/答/题"落在 45pt 带外，于是"请勿"被拼进最后一题题干末尾。
    判据必须看"这一列字本身"（同 x 带 ≥4 框、≥3 个不同字、贴边 100pt 内），
    不能为了抓它把 45pt 放宽——那会吃掉横版卷左半正文的行尾。
    """
    from app.extractor.lines import drop_margin_items
    col = [("弥", 850.3, 97.9), ("封", 850.3, 159.1), ("湾", 838.1, 167.8),
           ("线", 850.3, 221.0), ("内", 852.5, 283.7), ("请", 853.9, 345.6),
           ("勿", 852.5, 405.4), ("弥", 839.5, 462.2), ("答", 851.0, 466.6),
           ("题", 852.5, 528.5)]
    items = [{"text": t, "x": x, "y": y, "w": 9, "h": 12} for t, x, y in col]
    items.append({"text": "6.你做过什么手工作品？", "x": 70, "y": 300, "w": 400, "h": 14})
    keep, dropped = drop_margin_items(items, 900.0)
    assert len(dropped) == 10, f"密封线没摘干净：{[(d['x'], d['text']) for d in dropped]}"
    assert [k["text"] for k in keep] == ["6.你做过什么手工作品？"], keep
    # 反例一：同样的字在页面中间（答题格里的"答"）不许丢
    mid = [{"text": t, "x": 400, "y": 100 + i * 30, "w": 9, "h": 12}
           for i, (t, _x, _y) in enumerate(col)]
    _k, dr = drop_margin_items(mid, 900.0)
    assert not dr, f"页面中间的竖排答案字被当密封线吃了：{[d['text'] for d in dr]}"
    # 反例二：贴边但只有 3 个字（普通竖排短词）不许成列丢弃
    few = [{"text": t, "x": 850, "y": 100 + i * 60, "w": 9, "h": 12}
           for i, t in enumerate(["请", "勿", "答"])]
    _k, dr = drop_margin_items(few, 900.0)
    assert not dr, f"三个字不够成列，误杀了：{[d['text'] for d in dr]}"


def photo_route_uses_real_page_size_and_margin_filter():
    """接线检查：`lines_from_image` 必须按【图片尺寸】建页、并真的调用页边过滤。

    以前照片路线直接用 OCR 框的最大右缘当页宽（左侧正文只到 ~840pt），
    45pt 页边规则算在 840 上 → 密封线 x≈850 全落在"页面中间"，一个都不丢。
    这里把 OCR 换成固定条目（模具图用 PIL 真造 1700×1190），只测接线：
    页宽必须 = 1224pt，noise 必须带出 8 个框，正文不许少。
    """
    import os
    import tempfile
    from PIL import Image
    from app.extractor import lines as L
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "照片.png")
        # ⚠️ 长边必须 ≥ `ocr_image_auto` 的放大阈值（1600px）：否则模具图会被
        # 先放大 2 倍 OCR、坐标再除回原图，假 OCR 的 x 也会被除半 ——
        # 950px 的密封线跑到页面中间，这条接线测试就测不到页边过滤了。
        Image.new("RGB", (1700, 1190), (255, 255, 255)).save(path)
        fake = [{"text": "6.你做过什么手工作品？", "x": 120, "y": 500,
                 "w": 700, "h": 24, "conf": 0.95}]
        for i, t in enumerate("密封线内请勿答题"):
            fake.append({"text": t, "x": 1650, "y": 100 + i * 90,
                         "w": 16, "h": 24, "conf": 0.9})
        old = L.O.ocr_image
        L.O.ocr_image = lambda _p: [dict(x) for x in fake]
        try:
            r = L.lines_from_image(path)
        finally:
            L.O.ocr_image = old
        assert abs(r["pages"][0]["width_pt"] - 1224.0) < 0.5, \
            f"页宽没按图片尺寸算：{r['pages'][0]['width_pt']}"
        assert abs(r["pages"][0]["height_pt"] - 856.8) < 0.5, r["pages"][0]
        assert len(r.get("noise") or []) == 8, r.get("noise")
        assert r["pages"][0].get("margin_dropped") == 8, r["pages"][0]
        assert r["lines"] == ["6.你做过什么手工作品？"], r["lines"]


t("装订栏不许混进题干", margin_strip_never_enters_body_text)
t("贴边正文行尾不许被当装订栏（反例）", margin_filter_must_not_eat_body_line_ends)
t("照片边缘密封线整列要摘掉（含两个反例）", sealing_column_on_photo_edge_is_dropped)
t("照片路线用真实页宽且接上页边过滤", photo_route_uses_real_page_size_and_margin_filter)


def scan_page_columns_verified_with_ink():
    """扫描页的栏界要用图片墨迹复核：单栏卷的图形投影会切出假栏。

    新下载 P1（DOCX 里贴 6 张整页扫描图）实测：单栏页被切成 2–3 栏，
    同一行的 A/B 选项分进不同栏，题干按"栏优先"排出来整段错位
    （Q4 混入 Q7 的选项、Q18 落到下一大题、Q21 被拆成两半）。
    """
    import os
    import tempfile
    from PIL import Image, ImageDraw
    from app.extractor import lines as L

    with tempfile.TemporaryDirectory() as td:
        # ⚠️ 长边 ≥1600px，避开 `ocr_image_auto` 的 2 倍放大：放大后坐标会被
        # 除半，假 OCR 的两栏条目会挤进同一栏，这条测试就测不到墨迹复核了。
        full = os.path.join(td, "full.png")
        im = Image.new("RGB", (1600, 2400), "white")
        d = ImageDraw.Draw(im)
        for y in range(240, 2240, 160):
            d.line([(80, y), (1520, y)], fill="black", width=12)
        im.save(full)
        bands = [{"x0": 0.0, "x1": 552.0}, {"x0": 552.0, "x1": 1152.0}]
        got = L.verify_bands_with_image(full, bands, 1152.0)
        assert len(got) == 1, f"单栏页的假栏缝没被墨迹否掉：{got}"

        # 反例：真两栏（中间有贯通空白）不许被收成一栏
        two = os.path.join(td, "two.png")
        im2 = Image.new("RGB", (1600, 2400), "white")
        d2 = ImageDraw.Draw(im2)
        for y in range(240, 2240, 160):
            d2.line([(80, y), (640, y)], fill="black", width=12)
            d2.line([(960, y), (1520, y)], fill="black", width=12)
        im2.save(two)
        got2 = L.verify_bands_with_image(two, bands, 1152.0)
        assert len(got2) == 2, f"真栏缝被墨迹吃掉了：{got2}"

        # 接线：假 OCR 条目 + 全幅墨迹，阅读顺序必须按行（A、B 同一行）
        fake = [{"text": t, "x": x, "y": 200 + i * 240, "w": 80, "h": 56, "conf": 0.9}
                for i in range(8) for x, t in ((320, "A"), (880, "B"))]
        old_ocr, old_detect = L.O.ocr_image, L.detect_columns
        L.O.ocr_image = lambda _p: [dict(x) for x in fake]
        L.detect_columns = lambda _items, _w: [dict(b) for b in bands]
        try:
            r = L.lines_from_image(full)
        finally:
            L.O.ocr_image, L.detect_columns = old_ocr, old_detect
        assert r["pages"][0]["columns"] == 1, \
            f"扫描图路线没接墨迹复核：{r['pages'][0]}"
        assert len(r["lines"]) == 8 and "A" in r["lines"][0] and "B" in r["lines"][0], \
            f"阅读顺序还是栏优先：{r['lines'][:3]}"
t("扫描页栏界用图片墨迹复核（含真两栏反例）", scan_page_columns_verified_with_ink)


def gutter_must_be_wide_and_unblocked():
    """栏缝的三条底线：不贴页边、够宽、没有文字行横穿。

    新下载 P2 第 3 页（竖式计算）与 P3 第 2 页（2-up 计算题）实测：
    只检查"有一条空白带"时，字间 1–2% 的空隙、上下两组题之间的空白、
    页边留白都会通过，把 ①②③④ 按列重新排序，②④ 被吞进下一大题。
    """
    import os
    import tempfile
    from PIL import Image, ImageDraw
    from app.extractor import lines as L

    with tempfile.TemporaryDirectory() as td:
        bands = [{"x0": 0.0, "x1": 144.0}, {"x0": 144.0, "x1": 288.0}]

        # ① 字间空隙：左右两段之间只有 6px（1.5% 页宽）的空白，不算栏缝
        hair = os.path.join(td, "hair.png")
        im = Image.new("RGB", (400, 600), "white")
        d = ImageDraw.Draw(im)
        for y in range(60, 560, 40):
            d.line([(20, y), (197, y)], fill="black", width=3)
            d.line([(203, y), (380, y)], fill="black", width=3)
        im.save(hair)
        got = L.verify_bands_with_image(hair, bands, 288.0)
        assert len(got) == 1, f"1.5% 宽的字间空隙被当成栏缝：{got}"

        # ② 2-up 排版：左右两题并排，但通栏标题横穿中间 → 必须按行读
        twoup = os.path.join(td, "twoup.png")
        im = Image.new("RGB", (400, 600), "white")
        d = ImageDraw.Draw(im)
        for y in range(50, 560, 150):       # 通栏标题/大题说明（细线，占高 <2%）
            d.line([(30, y), (370, y)], fill="black", width=1)
        for y in range(100, 560, 120):      # 左右并排的小题
            d.line([(40, y), (150, y)], fill="black", width=3)
            d.line([(250, y), (360, y)], fill="black", width=3)
        im.save(twoup)
        got2 = L.verify_bands_with_image(twoup, bands, 288.0)
        assert len(got2) == 1, f"2-up 排版被切成两栏，②④会排到①③后面：{got2}"

        # ③ 页边留白：候选靠近左边，但左边那条空白是版心外留白，不是栏缝
        margin = os.path.join(td, "margin.png")
        im = Image.new("RGB", (400, 600), "white")
        d = ImageDraw.Draw(im)
        for y in range(60, 560, 40):
            d.line([(80, y), (380, y)], fill="black", width=3)
        im.save(margin)
        near_left = [{"x0": 0.0, "x1": 30.0}, {"x0": 30.0, "x1": 288.0}]
        got3 = L.verify_bands_with_image(margin, near_left, 288.0)
        assert len(got3) == 1, f"页边留白被当成栏缝：{got3}"

        # 反例：真两栏（宽缝 + 无文字行横穿）必须保住
        two = os.path.join(td, "two.png")
        im = Image.new("RGB", (400, 600), "white")
        d = ImageDraw.Draw(im)
        for y in range(60, 560, 40):
            d.line([(20, y), (160, y)], fill="black", width=3)
            d.line([(240, y), (380, y)], fill="black", width=3)
        im.save(two)
        got4 = L.verify_bands_with_image(two, bands, 288.0)
        assert len(got4) == 2, f"真栏缝被三条底线误杀：{got4}"
t("栏缝要够宽且没有文字行横穿（含页边/2-up/真两栏反例）", gutter_must_be_wide_and_unblocked)


def broken_brackets_get_closed():
    """OCR 读丢半个括号的填空必须补回来 —— 括号不配对的题【做不了】。

    真实卷：综合实践单元.pdf 一、填空第 1、2 题，原文"一端指（ ），一端指（ ）。"
    OCR 只给到"…一端指 ),一端指("。
    """
    from app.extractor.rules import repair_blanks
    t1, n1 = repair_blanks("1.指南针是我国古代四大发明之一，它的特点是一端指 ),一端指(")
    assert n1 >= 1 and t1.count("（") + t1.count("(") == t1.count("）") + t1.count(")"), \
        f"没补配对：{t1} / {n1}"
    t2, n2 = repair_blanks("2.傍晚，爸爸欣赏落日时，他的前面是（       )，后面是（")
    assert n2 == 1 and t2.rstrip().endswith("）"), f"行尾括号没补：{t2}"
    assert t2.count("（") + t2.count("(") == t2.count("）") + t2.count(")")


def broken_bracket_repair_leaves_real_text_alone():
    """反例：选项标号、序号、已配对的空位、算式冒号，一个都不许改。"""
    from app.extractor.rules import repair_blanks
    for s in ["A）对 B）错", "（1）小明今年8岁，（2）爸爸今年（ ）岁。",
              "选择题：4×7 读作(    )", "3:4=（ ）：12",
              "5.同学们将面向（       )站立。", "一、填空。(每空1分,共16 分)"]:
        out, n = repair_blanks(s)
        assert n == 0 and out == s, f"误伤正文：{s!r} → {out!r}（补了 {n} 处）"


t("读残的填空括号要补全", broken_brackets_get_closed)
t("补括号不许误伤序号/选项/已配对空位（反例）", broken_bracket_repair_leaves_real_text_alone)


def ocr_bracket_tail_digit_is_repaired():
    """P4 一.1：OCR 把句尾「）。」读成「3.」——第二空消失，孩子只能填一个空。

    证据：P4 证据包 06_questions.json 里 q1 blanks=1、题干收尾「读作（」；
    真实卷该句是「读作（            ）。」。
    """
    from app.extractor.rules import repair_blanks
    out, n = repair_blanks("读作（            3.")
    assert n == 1, f"没认出被读成数字的右括号：{out!r}（补了 {n} 处）"
    assert out.rstrip() == "读作（    ）.", f"补出来的空位不对：{out!r}"
    # 反例：没有空白的小问编号/正文数字一个都不许动
    for s in ["（3.求面积", "（3.）", "一共（            3）个", "（12分）"]:
        out2, n2 = repair_blanks(s)
        assert n2 == 0 and out2 == s, f"误伤：{s!r} → {out2!r}（补了 {n2} 处）"


def circled_option_numbers_are_repaired():
    """P4 三.2：OCR 把圈号 ② 读成数字 2 —— 选项只剩①②两条，孩子没法选。"""
    from app.extractor.rules import _circled_options
    opts = _circled_options("①缩小20倍         2扩大20倍         ③不变")
    assert len(opts) == 3, f"选项还是只有 {len(opts)} 条：{opts}"
    assert opts[1].startswith("②") and "扩大20倍" in opts[1], opts
    # 反例一：真正的数字选项 26 不许被吃成「②6」
    got = _circled_options("①13         26         ③39")
    assert any("26" in o and not o.startswith("②") for o in got), f"把数字选项改坏了：{got}"
    # 反例二：圈号齐全时一个字都不许动
    s = "①东 ②南 ③西 ④北"
    assert _circled_options(s) == ["①东", "②南", "③西", "④北"], _circled_options(s)


t("句尾「）。」被 OCR 读成数字要补回空位（含小问编号反例）", ocr_bracket_tail_digit_is_repaired)
t("圈号被 OCR 读成数字要补回（含纯数字选项反例）", circled_option_numbers_are_repaired)


def dangling_tail_residue_is_stripped():
    """题干句尾的孤立右括号、粘上的选项字母要清掉（P5 Q12/Q16/Q4 实测）。

    真实卷：
      · 12 "…占整图的多少， ） ）" —— 图里两个空位的右括号被读进题干；
      · 16 "…拼成近似的长方形。 ）"；
      · 4  "…是（        ） B            D" —— A/C 选项是图，只捡到 B、D。
    反例：正常配对的空位、正文中间的半边括号不许动（后者交给待校对）。
    """
    from app.extractor.rules import strip_dangling_tail
    assert strip_dangling_tail("…占整图的多少， ） ）") == "…占整图的多少"
    assert strip_dangling_tail("…拼成近似的长方形。 ）") == "…拼成近似的长方形。"
    assert strip_dangling_tail("…的是（        ） B            D") == "…的是（        ）"
    for s in ["…填在（        ）里。", "…是（ ），后面是（ ）",
              "…后面是（ 夜晚，…", "…（取3.14）", "…是（        )"]:
        assert strip_dangling_tail(s) == s, f"误伤：{s!r} → {strip_dangling_tail(s)!r}"

    # 接线：extract 出来的题干必须已经清理过
    r = extract(["二、填空题。",
                 "12.观察下图，分别用不同的数表示图中阴影部分占整图的多少， ） ）",
                 "13.如下图所示，把一个圆分成若干等份后，可以拼成近似的长方形。 ）",
                 "14.下面图形中，不是轴对称图形的是（        ） B            D"])
    got = {q["q_no"]: q["text"] for q in qs(r)}
    assert got[12].endswith("多少"), got[12]
    assert got[13].endswith("长方形。"), got[13]
    assert got[14].endswith("（        ）"), got[14]
    assert "B" not in got[14] and "D" not in got[14], got[14]
t("句尾孤立的右括号/选项字母要清掉（含反例）", dangling_tail_residue_is_stripped)


def unbalanced_paren_after_repair_is_flagged_not_silently_showed():
    """补不回来的括号（缺在句子中间，位置猜不出来）必须标"待人工校对"。

    真实卷：综合实践单元.pdf 一、填空第 2 题
      "…他的前面是（ ），后面是（ 夜晚，爸爸面向北极星时…"
    第二个空位缺的"）；"在句子中间，硬补就是编造题干。
    标出来让校对页挂角标，比悄悄印一张读不通的卷子好。
    """
    r = extract(["一、填空。",
                 "2.傍晚，爸爸欣赏落日时，他的前面是（ ），后面是（ 夜晚",
                 "，他的后面是（ ）。", "3.第三题的题干文字。"])
    q = [x for x in qs(r) if "爸爸欣赏落日" in x["text"]]
    assert q, [x["text"] for x in qs(r)]
    assert q[0]["needs_review"], f"括号不配对却没标待校对：{q[0]['text']}"
    assert "括号" in (q[0].get("review_note") or ""), q[0].get("review_note")
    # 反例：配好对的题不许被这条拖累（题号会被剥掉，所以按 q_no 找）
    ok = [x for x in qs(r) if x.get("q_no") == 3]
    assert ok, [x.get("q_no") for x in qs(r)]
    assert not ok[0].get("review_note"), ok[0]


t("补不回的括号要标待校对（含反例）", unbalanced_paren_after_repair_is_flagged_not_silently_showed)


def wiring_margin_filter_actually_runs_on_a_page():
    """接线检查：装订栏过滤必须真的挂在 `annotated_lines_from_pdf` 里。

    变异闸门打坏调用点（`items, margin = items, []`）时，上面那条直接调函数的
    单元测试【照样全绿】—— 因为它测的是函数会什么，不是"有没有人用它"。
    这条走完整链路：造一张带竖排装订栏的 PDF → 读行 → 断言正文里没有"学号/姓名"。
    """
    import os
    import tempfile
    import pymupdf
    from app.extractor.lines import annotated_lines_from_pdf
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "装订栏卷.pdf")
        doc = pymupdf.open()
        page = doc.new_page(width=841, height=653)
        page.insert_text((68, 60), "综合实践达标练习", fontsize=14, fontname="china-s")
        # 装订栏：竖排词，其中"学号："故意和正文同一基线 ——
        # 不同基线的话过滤关掉也只是多出行，测不出"混进题干"这个真症状。
        for w, wy in zip(["学校：", "班级：", "姓名：", "学号："], [140, 162, 184, 219]):
            page.insert_text((10, wy), w, fontsize=11, fontname="china-s")
        page.insert_text((64, 197), "一、填空。", fontsize=11, fontname="china-s")
        page.insert_text((64, 219), "1. 小朋友们在操场上辨认方向。", fontsize=11, fontname="china-s")
        doc.save(path)
        doc.close()
        r = annotated_lines_from_pdf(path)
        assert not any(("学号" in l or "姓名" in l or "学校" in l or "班级" in l)
                       for l in r["lines"]), f"装订栏字还是混进了正文：{r['lines'][:8]}"
        assert any("小朋友们在操场上辨认方向" in l for l in r["lines"]), r["lines"][:8]
        assert sum(p.get("margin_dropped", 0) for p in r["pages"]) >= 2, \
            f"过滤没跑：{[p.get('margin_dropped') for p in r['pages']]}"


def wiring_bracket_repair_actually_runs_in_extract():
    """接线检查：补括号必须真的在 extract() 的成题路径上（同上，防"函数对、线没接"）。"""
    r = extract(["一、填空。", "1.傍晚，爸爸欣赏落日时，他的前面是（       )，后面是（"])
    q = qs(r)
    assert q, [x["text"] for x in q]
    t = q[0]["text"]
    assert t.count("（") + t.count("(") == t.count("）") + t.count(")"), f"没补：{t}"
    assert q[0].get("blanks_closed"), q[0]


print()
print("[题号单独成行]")


def bare_number_line_opens_a_question():
    """题号被 OCR 切成单独一行、题干在整张图底下 —— 这道题不能凭空消失。

    真实卷：第三单元练习题 第 4 天「动物聚会」第 3 题，"3" 在 y=450（图顶），
    题干"观察图片，提出两个能用6×2解决的问题。"在 y=570（图底）。
    不认这个"3"，整卷就少一道题（审计器报"断号：缺第 3 题"）。
    """
    lines = [
        "一、看图回答问题。",
        "1. 想一想，连一连。",
        "3×4  4×3  4个3  3个4",
        "2",
        "观察图片，提出两个能用6×2解决的问题。",
        "问题一：      问题二：",
        "3. [思维探究]一共有多少个？把你的思考方法在图上圈一圈。",
    ]
    q = qs(extract(lines))
    nos = [x.get("q_no") for x in q]
    assert nos == [1, 2, 3], f"题号序列断了（少一道题孩子就少做一道）：{nos}"
    assert q[1]["text"].startswith("观察图片"), q[1]["text"]


def bare_number_must_be_the_next_in_sequence():
    """反例：数字不接着上一题、连着排的空数字 —— 都不许开题。"""
    # ① "5" 不接在 1 后面（图里的孤立数字、连线项）→ 并进题干，不开新题
    q = qs(extract(["一、填空。", "1. 第一题的题干文字内容。", "5", "第二行正文内容。"]))
    assert [x.get("q_no") for x in q] == [1], f"图里的孤立数字被当成了题号：{q}"
    assert "5" in q[0]["text"], q[0]["text"]
    # ② 连着排的数字（统计图上的 1 2 3 刻度）→ 只开一道，不许切出三道空题
    q2 = qs(extract(["一、看图填空。", "1. 第一题的题干文字内容。", "2", "3", "4"]))
    assert [x.get("q_no") for x in q2] == [1, 2],         f"一串刻度数字被切成好几道空题：{[x.get('q_no') for x in q2]}"
    # ③ 大题开头就冒出一个数字（没有上一题可接）→ 不开
    q3 = qs(extract(["二、解决问题。", "3", "这一行才是真正的题干内容。"]))
    assert [x.get("q_no") for x in q3] == [None], f"没有上一题时凭空认了题号：{q3}"


def wiring_bare_number_actually_runs_in_extract():
    """接线检查：判据必须真的接在 extract() 的成题路径上（防"函数对、线没接"）。"""
    from app.extractor.rules import _bare_number_starts_question as f
    prev = {"text": "想一想，连一连。"}
    assert f("3", 2, prev, 2) == 3, "判据本身就没通过（那上面两条测的是空气）"
    assert f("3", 2, {"text": ""}, 2) is None, "上一题还没有题干时不该认"
    r = extract(["一、看图回答问题。", "1. 第一题的题干文字内容。", "2",
                 "第二题的题干在图底下。"])
    assert [x.get("q_no") for x in qs(r)] == [1, 2],         f"判据算得出 2、成题路径却没开这道题（线没接）：{[x.get('q_no') for x in qs(r)]}"


print()
print("[区间起点：题号在图顶上]")


def span_anchor_lifts_to_the_bare_number_row():
    """裁图区间的起点要抬到那一行【题号】上，否则整张图落进上一题的区间。"""
    from app.extractor.lines import spans_for_questions
    ann = [
        {"text": "1. 想一想，连一连。", "x": 60, "y": 319, "w": 128, "h": 12,
         "page": 1, "col": 0},
        {"text": "2", "x": 60, "y": 450, "w": 11, "h": 12, "page": 1, "col": 0},
        {"text": "观察图片，提出两个能用6×2解决的问题。", "x": 75, "y": 570,
         "w": 294, "h": 12, "page": 1, "col": 0},
    ]
    questions = [
        {"q_no": 1, "text": "想一想，连一连。", "source_lines": ["想一想，连一连。"]},
        {"q_no": 2, "text": "观察图片，提出两个能用6×2解决的问题。",
         "source_lines": ["观察图片，提出两个能用6×2解决的问题。"]},
    ]
    pages = [{"page": 1, "width_pt": 595, "height_pt": 842,
              "bands": [{"x0": 0.0, "x1": 595.0}]}]
    sp = {int(s["q_index"]): s for s in spans_for_questions(ann, questions, pages=pages)}
    assert abs(sp[1]["y0"] - 450) < 1, f"区间起点没抬到题号行：{sp[1]['y0']}"
    assert abs(sp[0]["y1"] - 450) < 1, f"上一题的区间下界没收到题号行：{sp[0]['y1']}"

    # 反例：那一行的数字和本题题号对不上 → 不许抬（否则区间会吞掉上一题的图）
    ann2 = [dict(ann[0]), {**ann[1], "text": "7"}, dict(ann[2])]
    sp2 = {int(s["q_index"]): s for s in spans_for_questions(ann2, questions, pages=pages)}
    assert abs(sp2[1]["y0"] - 570) < 1, f"数字对不上也把区间往上抬了：{sp2[1]['y0']}"


t("题号单独成行也要开一道题", bare_number_line_opens_a_question)
t("反例：不接号的数字不许开题", bare_number_must_be_the_next_in_sequence)
t("裸题号判据真的接在成题链路上", wiring_bare_number_actually_runs_in_extract)
t("区间起点抬到图顶上那行题号（含反例）", span_anchor_lifts_to_the_bare_number_row)

def header_guard_skips_numbered_lines():
    """带题号的行永远不是"卷首信息行"，哪怕句子里有"答案/得分"这些表头词。

    真实坑：合成夹具里"1. 看图填方向，把答案写在括号里。"被 HEADER_HINTS 的
    "答案"命中，而这条判据跑在【第一道带题号的题】之前（seen_numbered 还是 False），
    于是整道第 1 题被当表头丢掉 —— 孩子看到的就是卷子少一题。
    """
    from app.extractor.rules import _looks_like_header as hdr
    # 正对照：真的卷首信息行还是要认
    assert hdr("班级＿＿＿＿ 姓名＿＿＿＿ 得分＿＿＿＿"), "卷首信息行不认了，表头会混进题目"
    assert hdr("满分：100分  时间：60分钟"), "卷首信息行不认了"
    # 这条判据要挡的就是它
    assert not hdr("1. 看图填方向，把答案写在括号里。"), "带题号的行仍被当表头"
    q = qs(extract(["一、看图填空。", "1. 看图填方向，把答案写在括号里。",
                    "2. 第二题的题干文字内容。"]))
    assert [x.get("q_no") for x in q] == [1, 2], \
        f"第 1 题被当卷首信息行丢了：{[x.get('q_no') for x in q]}"


def subquestion_paren_markers_merge_into_their_question():
    """"1. …再完成练习。" 下面的（1）（2）是【小问】，不许各开一道题。

    真实卷：第三单元达标练习题A 第三大题 —— 旧行为切出"第 1 题、第 1 题、第 2 题、
    第 2 题"，卷面上凭空多出两道假题（审计器报"重号 2"）。
    """
    lines = ["三、分类积累词语。（15分）",
             "1. 先把本单元学习的词语补充完整，再完成练习。（12分）",
             "（1）以上词语都可以用来描写春天的景象。（2分）",
             "（2）选词填空：春天来了，满园的鲜花让人眼前一亮。（3分）",
             "2. 下面的词语中，分类最合理的是哪一组？（3分）"]
    q = qs(extract(lines))
    assert [x.get("q_no") for x in q] == [1, 2], \
        f"小问（1）（2）被当成新题，出现重号：{[x.get('q_no') for x in q]}"
    assert "以上词语都可以用来描写" in q[0]["text"], q[0]["text"]
    assert "选词填空" in q[0]["text"], q[0]["text"]
    # 小问号必须保留并各起一行：吞掉编号后四问会糊成一句，孩子分不清答哪问
    # （独立复测 report_A P1-05：Q23 的 (1)…(2)…(3)…(4)… 全部消失）。
    assert "（1）" in q[0]["text"] and "（2）" in q[0]["text"], q[0]["text"]
    assert "\n" in q[0]["text"], f"小问被拼成一段：{q[0]['text']!r}"
    subs = split_subqs(q[0]["text"])
    assert [s["no"] for s in subs] == [1, 2], f"小问切不出来：{subs}"

    # 反例：整段都用（1）（2）平排编号的卷，它们就是各一道题，不许并
    q2 = qs(extract(["二、选一选。（10分）",
                     "（1）下列词语中不属于同一类的是（  ）",
                     "（2）下面说法正确的是（  ）"]))
    assert [x.get("q_no") for x in q2] == [1, 2], \
        f"平排的（1）（2）被并成一道题了：{[x.get('q_no') for x in q2]}"


t("小问（1）（2）不许切成两道重号题（含反例）", subquestion_paren_markers_merge_into_their_question)

t("卷首信息行不许吃掉带题号的第一题", header_guard_skips_numbered_lines)
def ocr_temp_files_must_not_collide_across_runs():
    """同一份扫描卷连跑两遍（先出卷、再识别题型），OCR 的临时图不许撞名。

    真实故障（2026-10-19 用户在网页上点"开始出卷"→ 识别题型失败：HTTP 500）：
    固定名 `lines_p2.png` + MuPDF 的 `pix.save()` 会【先删再写】；上一遍收尾的
    `os.remove` 因为 OCR 还压着句柄而失败（异常被咽掉），文件留在那儿，
    第二遍删不掉 → FzErrorSystem: Permission denied → 整份卷 500。
    修法：临时名带 pid+随机后缀（撞不上）+ 用 Python 直接覆盖写（不先删）。
    """
    import shutil
    import pymupdf
    from app.extractor import lines as L
    from app.extractor import ocr as O

    out = os.path.join(ROOT, "var", "_ocr_collide")
    shutil.rmtree(out, ignore_errors=True)
    os.makedirs(out, exist_ok=True)
    src = os.path.join(out, "纯扫描页.pdf")
    d = pymupdf.open()
    pg = d.new_page(width=595, height=200)
    pix = pymupdf.Pixmap(pymupdf.csGRAY, pg.get_pixmap(dpi=72))   # 只有像素、没有文字层
    pg.insert_image(pgm := pymupdf.Rect(0, 0, 595, 200), pixmap=pix)
    d.save(src)
    d.close()

    seen: list[str] = []
    held: list = []

    def fake_ocr(path: str):
        seen.append(os.path.basename(path))
        held.append(open(path, "rb"))        # 模拟 OCR 那边句柄还没放
        return [{"text": "1. 一、选择题", "x": 40, "y": 40, "w": 140, "h": 14, "conf": 0.9}]

    old = O.ocr_image
    try:
        O.ocr_image = fake_ocr
        r1 = L.annotated_lines_from_pdf(src, work_dir=os.path.join(out, "w1"), dpi=100)
        r2 = L.annotated_lines_from_pdf(src, work_dir=os.path.join(out, "w1"), dpi=100)
    finally:
        O.ocr_image = old
        for f in held:
            f.close()
        shutil.rmtree(out, ignore_errors=True)
    assert len(seen) == 2, f"两遍应该各渲染一次临时图：{seen}"
    assert len(set(seen)) == len(seen), f"临时图撞名（下一次就会去删一个被占用的文件）：{seen}"
    assert r2["lines"], "第二遍一行都没读到（说明这一页被静默跳过了）"
    assert not any(p["route"] in ("ocr-write-failed", "ocr-missing") for p in r2["pages"]), r2["pages"]


t("OCR 临时图唯一名：同一份卷跑两遍不许 500", ocr_temp_files_must_not_collide_across_runs)

t("装订栏过滤真的接在读页链路上", wiring_margin_filter_actually_runs_on_a_page)
t("补括号真的接在成题链路上", wiring_bracket_repair_actually_runs_in_extract)


def sub_questions_split_for_display_but_stay_one_question():
    """一道大题里的（1）（2）（3）要【分行分空作答】，但【不能拆成三道题】。

    用户拿真实卷问的："一个题目里面有三个小题，非得挤在一个框子里，怎么分别作答呢？"
    两头都要顾：拆成独立题会重排题号、丢分值（上一轮就是被这个搞出"重号假题"的）；
    完全不拆又看不清哪一空归哪一问。所以题号/分值/批改仍按一道题走，
    排版和作答按小问走 —— 这里测的是"小问怎么切"。
    """
    from app.extractor.rules import split_subqs

    t = ("(1)2只2只地数,数了(    )次,一共有(    )只蝴蝶。 "
         "(2)3只3只地数,数了(    )次,一共有(    )只蝴蝶。 "
         "(3)4只4只地数,数了(    )次,一共有(    )只蝴蝶。")
    subs = split_subqs(t)
    assert [s["no"] for s in subs] == [1, 2, 3], f"三问没切出来：{subs}"
    assert [s["blanks"] for s in subs] == [2, 2, 2], f"每问该有两个空：{subs}"
    assert sum(s["blanks"] for s in subs) == 6, "小问空位数要能对上做题页的 6 个输入框"

    # 反例①：填空的空位写作"（   ）"，没有数字 —— 绝不能当成小问号
    assert split_subqs("一年有（    ）个月，其中大月有（    ）个。") == [], \
        "把空位当成了小问号"
    # 反例②：只出现一次编号 → 不是一串小问
    assert split_subqs("（1）这一问后面再没有第二问。") == []
    # 反例③：已经填了数的括号（"3×4=（12）"）编号不从 1 起、也不递增
    assert split_subqs("算式 3×4=（12），5×6=（30）") == [], "把算式结果当成了小问号"
    # 反例④：编号跳号（1）、（3）—— 中间丢了（2），宁可整条不拆
    assert split_subqs("（1）第一问。（3）第三问。") == [], "跳号也照样拆，会把小问标错"
    # 反例⑤：光杆编号（内容在表格里，OCR 只回"（1）（2）（3）"）→ 拆出来是三行空白
    assert split_subqs("数一数。（1）（2）（3）") == [], "把没有正文的编号拆成三个空行"
    # 反例⑤b：编号之间有空格、都在断句边界上，但第一段是空的
    # ——边界判据拦不住它，得靠 seg 为空就整条不拆（否则打印出一行空白小问）
    assert split_subqs("（1） （2） 做一做") == [], "空小问被拆出来了"
    # 反例⑥：行内的普通数字（"第 1 组…第 2 组…"）—— 括号的【成对形状】才是小问号
    assert split_subqs("第 1 组有 5 个，第 2 组有 3 个，第 3 组有 7 个。") == [], \
        "括号可有可无时，量词数字被当成了小问号"

    # 引导语要留住：「看图填一填。（1）…（2）…」里的"看图填一填"不是小问，也不能丢
    lead = split_subqs("看图填一填。 （1）数是几个几？（2）一共有多少？")
    assert lead and lead[0].get("lead") == "看图填一填。", f"引导语丢了：{lead}"
    assert len(lead) == 2 and [x["no"] for x in lead] == [1, 2], lead

    # 反例⑦：班级编号不是小问号（P5 Q24「六（1）班、六（2）班」被拆成两段，
    # 第一段只剩"班、六"，打印出来题干错位）。真正的（1）（2）在句读后，照切。
    cls = split_subqs("张老师是六（1）班、六（2）班的体育老师。"
                      "（1）按成绩段整理数据。（2）回答下面的问题。")
    assert [x["no"] for x in cls] == [1, 2], f"班级编号被当小问号：{cls}"
    assert cls[0].get("lead") == "张老师是六（1）班、六（2）班的体育老师。", \
        f"引导语被班级编号截断：{cls}"


t("（1）（2）（3）小问切得出来也不误伤空位（含七个反例）",
  sub_questions_split_for_display_but_stay_one_question)


def long_blank_runs_count_as_blanks():
    """括号里【一整串空白】就是一个空位，不管几个空格。

    真卷实测（E:\\数学卷子\\第三单元练习题.pdf 第 1 页）：扫描 OCR 把空位读成
    "(      )" 六个空格，旧口径只认 1–4 个 → 卷面 33 个空位数出 3 个，
    做题页每道题只给一个输入框，孩子根本没法"分别作答"。
    """
    from app.extractor.rules import BLANK_RE

    real = "(1)2只2只地数,数了(      )次,一共有(      )只蝴蝶。"
    assert len(BLANK_RE.findall(real)) == 2, \
        f"真卷这一行的空位数不对：{len(BLANK_RE.findall(real))}（应为 2：两个 6 空格空位；(1) 是小问号不算）"
    assert len(BLANK_RE.findall("一年有（        ）个月，大月有（  ）个。")) == 2
    # OCR 常把一对括号读成"全角开 + 半角闭"的混配，只认成对写法就漏（真卷有 5 个）
    assert len(BLANK_RE.findall("数了（      )次,一共有(      ）只蝴蝶。")) == 2, \
        "混配括号的空位没数到"
    assert len(BLANK_RE.findall("算式：____________ ＝ ____________")) == 2
    # 反例①：括号里有内容 → 那是填了数的空/普通括号，不是空位
    assert len(BLANK_RE.findall("3×4=（12），5×6=（30）")) == 0, "把填了数的括号当成空位"
    # 反例②：括号里夹了字 → 不是空位
    assert len(BLANK_RE.findall("（ 见上图 ）不对，应该是（   ）")) == 1, "把带字的括号当成空位"
    # 反例③：一个下划线不是空位（卷面上的脚注线、破折号）
    assert len(BLANK_RE.findall("a_b")) == 0
t("括号里的长空格流要数成空位（含三个反例）", long_blank_runs_count_as_blanks)


def multiline_stems_close_their_blank_and_keep_subquestions_apart():
    """多行题干：行尾的空位要补齐，小问 (1)(2) 不许被空格拼成一句。

    真卷（周末作业(1).pdf）：第 1 页第 3 题 "…可以用3×4解决？（" 后面紧接
    一行选项条目，整段 repair_blanks 只看段尾 → 空位的右半边永远补不上；
    第 5 题 "…可以是（" 下一行就是 "(2)…"，成品印成
    "…可以是（ (2)▲+▲…"，两道小问糊成一句。
    """
    lines = [
        "1. 下面哪些问题可以用3×4解决？（",
        "牛奶  牛奶  牛奶  4元4元4元",
        "2. (1)■×7＞8+8+8+8+8，■表示的数可以是（",
        "(2)▲+▲+●=▲×5，●=（    ）个▲。",
    ]
    got = qs(extract(lines))
    assert got[0]["text"].startswith("下面哪些问题可以用3×4解决？（    ）"), got[0]["text"]
    assert "\n" in got[0]["text"], "条目行与题干没分开"
    assert got[1]["blanks"] >= 2, f"多行题干的空位少算了：{got[1].get('blanks')}"
    assert "\n(2)" in got[1]["text"], f"小问 (2) 没另起一行：{got[1]['text']}"
    assert got[1]["text"].count("（") == got[1]["text"].count("）"), got[1]["text"]
    assert not got[1].get("review_note"), got[1]
t("多行题干：行尾空位要补齐、小问 (1)(2) 不许拼成一句",
  multiline_stems_close_their_blank_and_keep_subquestions_apart)


def span_locates_the_number_glued_stem_row():
    """题号【粘在第一行行首】（"4.(1)一共有多少根胡萝卜?"）时，区间起点必须是这一行。

    真卷第 2 页实测：上一题的区间一路吃到 y=584.8，第 4 题的货架图和答题空
    整块落进第 3 题的裁框里 → 第 3 题的图混着别人的题，第 4 题一张图都没有。
    用户那句话就是这么来的："这题都对不上，怎么做啊"。
    """
    from app.extractor.lines import spans_for_questions
    ann = [
        {"text": "3.[自然科普]落叶腐(f)烂后会化作养分,促进植物生长。下面的树叶片数",
         "x": 41, "y": 407.0, "w": 480, "h": 19, "page": 2, "col": 0},
        {"text": "不能用2×4表示的是(       )。(填序号)",
         "x": 60, "y": 436.7, "w": 300, "h": 19, "page": 2, "col": 0},
        {"text": "CL            长长", "x": 60, "y": 482.9, "w": 400, "h": 18,
         "page": 2, "col": 0},
        {"text": "4.(1)一共有多少根胡萝卜?", "x": 41, "y": 519.5, "w": 210, "h": 19,
         "page": 2, "col": 0},
        {"text": "(2)5X2在图中表示什么意思？", "x": 60, "y": 584.8, "w": 220, "h": 19,
         "page": 2, "col": 0},
        {"text": "第三步提思维", "x": 41, "y": 701.9, "w": 100, "h": 19, "page": 2, "col": 0},
    ]
    questions = [
        {"q_no": 3, "text": "[自然科普]落叶腐(f)烂后会化作养分,促进植物生长。"
                           "下面的树叶片数 不能用2×4表示的是(       )。(填序号)",
         "source_lines": ["[自然科普]落叶腐(f)烂后会化作养分,促进植物生长。下面的树叶片数",
                          "不能用2×4表示的是(       )。(填序号)", "CL            长长"]},
        {"q_no": 4, "text": "(1)一共有多少根胡萝卜? (2)5X2在图中表示什么意思？",
         "source_lines": ["(1)一共有多少根胡萝卜?", "(2)5X2在图中表示什么意思？"]},
    ]
    pages = [{"page": 2, "width_pt": 595, "height_pt": 907,
              "bands": [{"x0": 0.0, "x1": 595.0}]}]
    sp = {int(s["q_index"]): s for s in spans_for_questions(ann, questions, pages=pages)}
    assert abs(sp[1]["y0"] - 519.5) < 1, \
        f"第 4 题的区间没从粘着题号的那一行开始：{sp[1]['y0']}（图会整个给上一题）"
    assert abs(sp[0]["y1"] - 519.5) < 1, \
        f"第 3 题的区间下界没收到第 4 题起手行：{sp[0]['y1']}"

    # 反例：行首的数字【不是本题题号】→ 不许拿这一行当起点（否则区间会乱吞）
    ann2 = [dict(x) for x in ann]
    ann2[3] = dict(ann[3], text="9.(1)一共有多少根胡萝卜?")
    sp2 = {int(s["q_index"]): s for s in spans_for_questions(ann2, questions, pages=pages)}
    assert abs(sp2[1]["y0"] - 584.8) < 1, \
        f"行首数字对不上也把区间往上抬了：{sp2[1]['y0']}"
t("题号粘在题干行首时，区间起点要落在这一行（含反例）",
  span_locates_the_number_glued_stem_row)


def span_locates_a_cross_page_merged_line():
    """跨页合并行要落到下一页页首那一行（含反例）。

    P2 Q4："4、估算"在页底、四个算式在下一页页顶，OCR 把两行并成
    "估算 901×8= …"。整条既不是任何 ann 行，也不是任何行的前缀/题号，
    locate 返回 -1 → 这道题没有区间，拿不到整题图；OCR 把 "÷" 读成 "+"
    也没法靠原图纠正。判据：页首那行是合并行的后缀时认它。
    """
    from app.extractor.lines import spans_for_questions
    ann = [
        {"text": "4、估算", "x": 34.2, "y": 528.8, "w": 46.8, "h": 15.1,
         "page": 3, "col": 0},
        {"text": "901×8=       98÷5≈   50÷7≈       989+605",
         "x": 25.9, "y": 8.3, "w": 260.3, "h": 13.7, "page": 4, "col": 0},
        {"text": "五、列式计算(5 分)", "x": 30, "y": 84.2, "w": 200, "h": 14,
         "page": 4, "col": 0},
    ]
    questions = [{"q_no": 4,
                  "text": "估算 901×8=       98÷5≈   50÷7≈       989+605",
                  "source_lines": ["估算 901×8=       98÷5≈   50÷7≈       989+605"]}]
    pages = [{"page": 3, "width_pt": 374.4, "height_pt": 589.0,
              "bands": [{"x0": 0.0, "x1": 374.4}]},
             {"page": 4, "width_pt": 374.4, "height_pt": 589.0,
              "bands": [{"x0": 0.0, "x1": 374.4}]}]
    spans = spans_for_questions(ann, questions, pages=pages,
                                sections=[{"name": "五、列式计算(5 分)"}])
    assert len(spans) == 1, f"跨页合并行没有区间：{spans}"
    sp = spans[0]
    assert int(sp["page"]) == 4 and abs(float(sp["y0"]) - 8.3) < 1, \
        f"区间没落在下一页页首：{sp}"
    assert abs(float(sp["y1"]) - 84.2) < 1, f"区间下界没收到下一道题：{sp}"

    # 反例：页首那行不是合并行的后缀（OCR 错得连内容都对不上）→ 不许硬认
    ann2 = [dict(a) for a in ann]
    ann2[1]["text"] = "801×8=       98÷5≈   50÷7≈       989+605"
    spans2 = spans_for_questions(ann2, questions, pages=pages)
    assert spans2 == [], f"对不上内容的后缀也被认成了区间：{spans2}"
t("跨页合并的题干要落到下一页页首（含反例）",
  span_locates_a_cross_page_merged_line)


def section_wall_uses_ocr_name_after_mark_repair():
    """大题标题补了标记后（打的→打▲的），硬墙锚点仍要落在原 OCR 行上。

    P6：改名后如果还拿新名字去找 ann，找不到标题 → 上一题的区间会越过
    页底漫进下一大题，整题截图把下一大题的题面也裁进来。
    """
    from app.extractor.lines import spans_for_questions
    ann = [
        {"text": "1. 285×9=", "x": 20, "y": 40, "w": 100, "h": 12, "page": 1, "col": 0},
        {"text": "四、笔算（打的要验算）（24%）", "x": 20, "y": 300, "w": 180,
         "h": 12, "page": 1, "col": 0},
        {"text": "2. 23÷7=", "x": 20, "y": 330, "w": 100, "h": 12, "page": 1, "col": 0},
    ]
    questions = [{"q_no": 1, "text": "285×9=", "source_lines": ["285×9="]},
                 {"q_no": 2, "text": "23÷7=", "source_lines": ["23÷7="]}]
    pages = [{"page": 1, "width_pt": 400, "height_pt": 550,
              "bands": [{"x0": 0, "x1": 400}]}]
    secs = [{"name": "四、笔算（打▲的要验算）（24%）",
             "ocr_name": "四、笔算（打的要验算）（24%）"}]
    sp = spans_for_questions(ann, questions, pages=pages, sections=secs)
    assert len(sp) == 2, f"区间数不对：{sp}"
    assert sp[0]["y1"] <= 300, f"第 1 题的区间越过了补过标记的大题标题：{sp[0]}"
    # 反例：没有 ocr_name（旧数据）时，新名字定位不到标题 → 上界会漫到下一题。
    # 这条反例证明「ocr_name 是硬墙还立着的唯一依据」，不是摆设。
    sp2 = spans_for_questions(ann, questions, pages=pages,
                              sections=[{"name": "四、笔算（打▲的要验算）（24%）"}])
    assert sp2[0]["y1"] > 300, f"没有 ocr_name 却仍找到了标题，反例失效：{sp2[0]}"
t("补过标记的大题标题仍是裁图硬墙（ocr_name 锚点）",
  section_wall_uses_ocr_name_after_mark_repair)


def span_expands_across_columns_for_cross_column_items():
    """一道题的 (2)(4) 印在右栏时，区间必须横跨两栏（含两个反例）。

    P5 第 5 页 Q17/Q18 实测：左栏是 (1)(3)，右栏是 (2)(4)，题锚点在左栏。
    旧区间只到栏界 135.3pt，整题截图把右栏两道小题切掉 —— 四道题只剩两道，
    孩子做不了。只认本题 source_lines 原样出现的行，隔壁题不许被框进来。
    """
    from app.extractor.lines import spans_for_questions
    ann = [
        {"text": "三、计算题（共2道小题）", "x": 46.8, "y": 25.2, "w": 73.4,
         "h": 8, "page": 5, "col": 0},
        {"text": "18.计算下面各题。", "x": 46.6, "y": 93.8, "w": 43.0,
         "h": 8, "page": 5, "col": 0},
        {"text": "(1)   4x3184", "x": 49.0, "y": 106.3, "w": 37.7,
         "h": 8, "page": 5, "col": 0},
        {"text": "(2)   279314497", "x": 149.3, "y": 107.8, "w": 44.4,
         "h": 8, "page": 5, "col": 1},
        {"text": "(3)         3", "x": 50.2, "y": 192.2, "w": 22.8,
         "h": 8, "page": 5, "col": 0},
        {"text": "(4)     25352571.7  x65", "x": 150.2, "y": 190.3, "w": 55.0,
         "h": 8, "page": 5, "col": 1},
        {"text": "四、解方程（共2道小题）", "x": 47.3, "y": 268.3, "w": 73.9,
         "h": 8, "page": 5, "col": 0},
        {"text": "20.", "x": 147.4, "y": 286.3, "w": 9.1, "h": 8,
         "page": 5, "col": 1},
    ]
    questions = [{"q_no": 18, "section": "三、计算题（共2道小题）",
                  "text": "计算下面各题。 (1)… (2)… (3)… (4)…",
                  "source_lines": ["计算下面各题。", "(1)   4x3184",
                                   "(2)   279314497", "(3)         3",
                                   "(4)     25352571.7  x65"]}]
    pages = [{"page": 5, "width_pt": 285.6, "height_pt": 404.16,
              "bands": [{"x0": 0.0, "x1": 135.31},
                        {"x0": 135.31, "x1": 285.6}]}]
    secs = [{"name": "三、计算题（共2道小题）"},
            {"name": "四、解方程（共2道小题）"}]
    sp = spans_for_questions(ann, questions, pages=pages, sections=secs)
    assert len(sp) == 1, sp
    assert float(sp[0]["x1"]) >= 195.0, \
        f"右栏的 (2)(4) 没被框进来：x1={sp[0]['x1']}"
    assert float(sp[0]["x0"]) <= 1.0, f"左栏起点被挪了：x0={sp[0]['x0']}"

    # 反例一：source_lines 里没有右栏行 → 区间必须留在本栏
    q1 = [dict(questions[0], source_lines=["计算下面各题。", "(1)   4x3184",
                                           "(3)         3"])]
    sp1 = spans_for_questions(ann, q1, pages=pages, sections=secs)
    assert float(sp1[0]["x1"]) <= 135.4, \
        f"没引用右栏也扩栏了：x1={sp1[0]['x1']}"

    # 反例二：右栏同文但 y 在区间之外（页脚）→ 不许把整栏扯进来
    ann2 = [dict(a) for a in ann]
    ann2.append({"text": "计算下面各题。", "x": 150.0, "y": 375.1, "w": 64.6,
                 "h": 8, "page": 5, "col": 1})
    sp2 = spans_for_questions(ann2, questions, pages=pages, sections=secs)
    assert float(sp2[0]["x1"]) <= 210.0, \
        f"区间外的右栏行把框拉长了：x1={sp2[0]['x1']}"
t("跨栏小题：区间要横跨两栏（含两个反例）",
  span_expands_across_columns_for_cross_column_items)


def glued_item_number_is_separated_from_decimal():
    """小题号粘在小数前面要拆开（含三个反例）。

    P3 第三大题第 1 小题：源卷是 "1. 14.06×3.5=  2. 4.25×6.8=  3. 12.5×2×0.8="，
    OCR 读成 "1.14.06×3.5= …"，看起来像 1.14×06×3.5。判据只认一个数里
    两个小数点且后面紧跟乘除；正常小数和日期不许动。
    """
    from app.extractor.rules import separate_glued_item_numbers
    raw = "1.14.06×3.5=            2.4.25×6.8=            3.12.5×2×0.8="
    got = separate_glued_item_numbers(raw)
    assert got.startswith("1、14.06×3.5="), f"第 1 小题没拆开：{got}"
    assert "2、4.25×6.8=" in got and "3、12.5×2×0.8=" in got, f"后两小题没拆开：{got}"

    # 反例①：正常小数（只有一个点）不许动
    assert separate_glued_item_numbers("3.5×2=7") == "3.5×2=7", "正常小数被改坏了"
    # 反例②：两个点但后面不是乘除（日期/版本号）不许动
    assert separate_glued_item_numbers("2024.10.01") == "2024.10.01", "日期被改坏了"
    # 反例③：两个点且是加法（不是小题号场景）不许动
    assert separate_glued_item_numbers("1.14.06+3.5") == "1.14.06+3.5", "加法式被误拆"
t("小题号粘在小数前要拆开（含三个反例）",
  glued_item_number_is_separated_from_decimal)


def stage_banner_is_furniture_not_a_figure():
    """教辅的【栏目横幅】（"第二步 练变式"）是页面家具，不许被当成配图裁下来。

    真卷实测（第三单元练习题 第 3 页第 1 题）：整块橙底反白的"第二步练变式"
    被裁成一张"配图"挂在题目下面，卷子上就是一张和题目无关的栏目名。
    """
    from app.extractor.rules import is_furniture_text
    assert is_furniture_text("第二步练变式"), "栏目横幅没当家具（会被裁成配图）"
    assert is_furniture_text("第三步 提思维")
    assert is_furniture_text("知识演练场")
    # 反例：题干里出现同一个关键词，不算横幅
    assert not is_furniture_text("3.【能力提升】看图列式，图中的圆点有几个？"), \
        "带题号的题干被当横幅抹掉了"
    assert not is_furniture_text("第一步，请把 3 个 4 改写成乘法算式，并读出算式。"), \
        "长句子被当横幅抹掉了"
t("教辅栏目横幅是家具不是配图（含反例）", stage_banner_is_furniture_not_a_figure)


def stage_banner_is_not_glued_to_the_question():
    """闯关横幅"第二关一方向应用我能行"不是题干，也不许接到上一题末尾。

    真实卷（综合实践单元(1) p1 第 5 题）：横幅在图上/下一段之前，OCR 单独一行，
    解析时被并进第 5 题题干。反例：句子里真的在说"第几关"的题干不许动。
    """
    from app.extractor.rules import _is_noise, is_furniture_text
    assert _is_noise("第二关一方向应用我能行"), "关卡横幅没当家具"
    assert is_furniture_text("第二关一方向应用我能行")
    assert not _is_noise("第二关我们来到了动物园"), "正常题干被当横幅抹掉了"
    assert not _is_noise("第二关的方向知识真有趣"), "正常题干被当横幅抹掉了"
    got = qs(extract(["5.小君想先参观书法馆，接下来需要向（   ）走。",
                      "第二关一方向应用我能行",
                      "三、公园畅游指南。（共21分）"]))
    assert got and "第二关" not in str(got[0].get("text")), \
        f"关卡横幅接进题干了：{[q.get('text') for q in got]}"
t("闯关横幅不许接在题干末尾（含反例）", stage_banner_is_not_glued_to_the_question)


def map_labels_before_the_number_do_not_hide_a_question():
    """行首是地图标签、题号在后面：仍要切开，且题干不许带题号（含两个反例）。"""
    got = qs(extract([
        "三、公园畅游指南。（共21分）",
        "东北           1.喷泉的西面是（      )，东面",
        "喷泉      游乐场            是（    )。(4分)",
        "莲池      假山            2.莲池在凉亭的（     )面，在",
        "凉亭            假山的（     )面。(4分)",
    ]))
    assert [q.get("q_no") for q in got] == [1, 2], \
        f"行首标签挡住了题号：{[(q.get('q_no'), q.get('text')) for q in got]}"
    assert got[0]["text"].startswith("喷泉的西面"), f"题干起手不对：{got[0]['text']!r}"
    assert not got[0]["text"].startswith("1."), f"题号留在题干里：{got[0]['text']!r}"
    assert not got[1]["text"].startswith("2."), f"题号留在题干里：{got[1]['text']!r}"
    # 反例①：序号跳号（上一题是 1，这行写 5.）→ 不认成新题
    skip = qs(extract(["三、公园畅游指南。（共21分）",
                       "东北           1.喷泉的西面是（      )，东面",
                       "莲池            5.莲池在凉亭的（     )面，在"]))
    assert [q.get("q_no") for q in skip] == [1], f"跳号被当新题：{skip}"
    # 反例②：行首是长正文（超过 8 字）再跟数字 → 不认
    longline = qs(extract(["三、公园畅游指南。（共21分）",
                           "1.喷泉的西面是（      )，东面",
                           "下面是一张公园平面示意图请看图           2.莲池在凉亭的（     )面"]))
    assert [q.get("q_no") for q in longline] == [1], f"长正文被当标签：{longline}"
    assert "下面是一张公园平面示意图请看图" in longline[0]["text"], "正文被吞了"


t("行首地图标签不许挡住题号（含两个反例）", map_labels_before_the_number_do_not_hide_a_question)


def blank_parentheses_are_not_item_gaps():
    """括号里的空白是【空位】，不是连线卡片的间隔。

    真实卷（周末作业(1).pdf 第 1 页第 1 题"想一想，连一连"）的卡片行
    `6+6   6个2   6×2 …` 必须保留两格以上空白；可同一份卷的
    `数了(      )次,一共有(      )只蝴蝶。` 按两格空白切也会切出三个"短条目"，
    旧版把整行判成卡片行：Python 侧给题干续行多插一个换行，网页端把空位里
    那 4 个空格原样留下（tests/subqs_assert.js 的"题干分行"当场红）。
    """
    from app.extractor.rules import _is_item_row
    assert _is_item_row("6+6   6个2   6×2   2×6   2+2+2+2+2+2   6+2   2个6"), \
        "真的连线卡片行没认出来"
    assert not _is_item_row("数了(      )次,一共有(      )只蝴蝶。"), \
        "空位里的空白被当成了条目间隔"
    assert not _is_item_row("一年有（        ）个月，大月有（  ）个。"), \
        "全角括号空位被当成了条目间隔"
    got = qs(extract(["1. 看图填一填。", "看数填空。",
                      "数了(      )次,一共有(      )只蝴蝶。"]))
    assert got and "\n" not in got[0]["text"], \
        f"题干被多插了换行（网页端会连空位一起变形）：{got[0]['text']!r}"


t("括号里的空白不是连线卡片的间隔（含反例）", blank_parentheses_are_not_item_gaps)


def worksheet_furniture_never_joins_the_stem():
    """教辅活页的出版社品牌行 / 填写栏 / 提示栏不是题干，也不许并进上一题。

    真实卷（周末作业(1).pdf，王朝霞《少而精》活页）：
      · p3 页脚 "王朝霞  少而精·活页同步练习册" 与
        "去除封面化繁为简品质优良价格公道" 两行都超过 24 字闸门，旧版直接
        并进了第 2 题的题干；
      · 同页 "班级/姓名  日期" 填写栏也被当成正文；
      · p4 "解题妙招"（楼数=层数+1）和 "知识打卡"（年龄科普）整段并进第 5 题。
    标题行上的家具（答案P122 / AI伴学 / 拍照秒批 / 页边装饰箭头）必须摘掉，
    只留下课时名 —— 否则分段名会带着广告印到卷面上。
    """
    from app.extractor.rules import _clean_title, _is_noise
    assert _is_noise("班级/姓名        日期"), "填写栏没被当家具"
    assert _is_noise("王朝霞            少而精·活页同步练习册"), "品牌行没被当家具"
    assert not _is_noise("小明做口算题，实际用时比计划少了3分钟。"), \
        "正常题干（带「实际用时」）被当页眉家具抹掉了"
    lines = ["第4课时 动物聚会   答案P122",
             "王朝霞            少而精·活页同步练习册",
             "去除封面化繁为简品质优良价格公道",
             "班级/姓名        日期",
             "1.第一题（    ）的题干内容。",
             "解题妙招",
             "楼数=层数+1",
             "2.第二题（    ）的题干内容。"]
    got = qs(extract(lines))
    assert [q.get("q_no") for q in got] == [1, 2], \
        f"家具行另立了一道假题：{[q.get('q_no') for q in got]}"
    text = " ".join(str(q.get("text") or "") for q in got)
    for bad in ("王朝霞", "活页同步练习册", "去除封面", "班级", "日期",
                "楼数", "层数"):
        assert bad not in text, f"家具内容「{bad}」混进题干：{text!r}"
    assert "第一题" in got[0]["text"] and "第二题" in got[1]["text"], \
        f"家具过滤把正常题干一起吞了：{[q.get('text') for q in got]}"

    # 标题行的三种真实家具各摘一次；反例：摘完不能只剩空串
    assert _clean_title("第4课时 动物聚会   答案P122") == "第4课时 动物聚会"
    assert _clean_title("第5课时快乐的动物   AI伴学") == "第5课时快乐的动物"
    assert _clean_title("第2天 画校园路线图[课时2]   拍照秒批") == "第2天 画校园路线图[课时2]"
    assert _clean_title("第3天 认识小数 <<<9") == "第3天 认识小数"


t("教辅活页的品牌行/填写栏/提示栏不当题干（含反例）",
  worksheet_furniture_never_joins_the_stem)


def colored_badges_are_found_but_decorations_are_not():
    """教辅活页的白字红底题号：OCR 检测网络整块漏掉，必须补检。

    真实卷（周末作业(1).pdf）：第 1 页 1/2/4 题、第 2 页 1/3/4 题的题号都是
    白字压粉色圆角块，RapidOCR 一个都没检出 → 整页并成一道无题号假题，
    用户看到的就是"连线条目被并段"。补检判据：①尺寸像题号（页宽 1%–5%）；
    ②右边紧跟一行汉字题干；③白字只识别、不检测；④OCR 已把它读成
    "2看一看"时把题号切开。
    """
    import numpy as np
    import cv2
    from app.extractor.badges import find_badges

    W, H = 1200, 300

    def page():
        return np.full((H, W, 3), 255, dtype=np.uint8)

    def pink(img, x, y, w, h):
        # BGR(60,60,220) → HSV 色相 0（红/粉那一档），和教辅徽标同色系
        cv2.rectangle(img, (x, y), (x + w, y + h), (60, 60, 220), -1)

    def white(img, x, y, w, h):
        cv2.rectangle(img, (x, y), (x + w, y + h), (255, 255, 255), -1)

    stem = {"text": "看一看，连一连。", "x": 70.0, "y": 45.0,
            "w": 260.0, "h": 20.0}

    def rec(text, conf=0.99):
        return lambda im: (text, conf)

    g = page(); pink(g, 20, 40, 40, 40); white(g, 32, 50, 16, 20)
    got = find_badges(g, [dict(stem)], recognize=rec("2"))
    assert len(got) == 1 and got[0]["text"] == "2.", f"白字红底题号没认出来：{got}"
    assert got[0]["badge"] is True, f"补检出的题号没打标记：{got[0]}"

    # 两位数：两个白字块要并成一个数字
    g = page(); pink(g, 20, 40, 40, 40)
    white(g, 26, 50, 10, 20); white(g, 40, 50, 10, 20)
    got = find_badges(g, [dict(stem)], recognize=rec("10"))
    assert len(got) == 1 and got[0]["text"] == "10.", f"两位数题号没认出来：{got}"

    # 反例①：徽标右边没有题干（装饰图标）→ 不许当题号
    g = page(); pink(g, 20, 40, 40, 40); white(g, 32, 50, 16, 20)
    assert find_badges(g, [], recognize=rec("2")) == [], "装饰图标被当题号"
    # 反例②：右边是英文而不是汉字 → 不是题干
    assert find_badges(g, [dict(stem, text="abc")], recognize=rec("2")) == [], \
        "英文标签被当题干"
    # 反例③：只识别通道读出来不是 1–2 位数字
    assert find_badges(g, [dict(stem)], recognize=rec("AB")) == [], "乱码被当题号"
    # 反例④：置信度 <0.5
    assert find_badges(g, [dict(stem)], recognize=rec("2", 0.3)) == [], \
        "低置信度识别被当题号"
    # 反例⑤：10% 页宽的大红方块（装饰图标/色块）不是题号。
    #   ⚠️ 别用长条横幅当反例：它先被长宽比挡住，尺寸上限那条判据等于没测。
    g = page(); pink(g, 20, 100, 120, 120); white(g, 60, 130, 40, 60)
    big_stem = dict(stem, x=160.0, y=150.0)
    assert find_badges(g, [big_stem], recognize=rec("5")) == [], "大色块被当题号"

    # 已有 OCR 框：OCR 把"2看一看"连成一条 → 切成"2.看一看"，不重复注入
    g = page(); pink(g, 20, 40, 40, 40); white(g, 32, 50, 16, 20)
    it = {"text": "2看一看", "x": 18.0, "y": 38.0, "w": 60.0, "h": 44.0}
    got = find_badges(g, [it], recognize=rec("2"))
    assert got == [] and it["text"] == "2.看一看", f"题号和题干没切开：{it}"
    # 反例⑥：压着徽标的是「第4课时」，不是数字 → 原样不动
    it2 = {"text": "第4课时 动物聚会", "x": 18.0, "y": 38.0,
           "w": 120.0, "h": 44.0}
    find_badges(g, [it2], recognize=rec("2"))
    assert it2["text"] == "第4课时 动物聚会", f"课时标题被改了：{it2}"
    # 反例⑦：OCR 把「第5课时」拆成「第」「5」「课时」三个框 —— 5 是课时序号，
    #   补成「5.」会让成品印出「第 5. 课时」（周末作业第 2 页真实发生）。
    it3 = {"text": "第", "x": 0.0, "y": 44.0, "w": 18.0, "h": 28.0}
    it4 = {"text": "5", "x": 20.0, "y": 38.0, "w": 40.0, "h": 44.0}
    it5 = {"text": "课时", "x": 64.0, "y": 44.0, "w": 40.0, "h": 28.0}
    got = find_badges(g, [it3, it4, it5], recognize=rec("5"))
    assert got == [] and it4["text"] == "5", f"课时序号被当题号：{it4}"


t("彩色题号徽标要认出来、装饰图标不许当题号（含七个反例）",
  colored_badges_are_found_but_decorations_are_not)


def page_header_ads_are_not_questions():
    """页眉的扫码批改/AI 伴学广告不是题目，更不许拼成一道无题号假题。

    真实卷（周末作业(1).pdf 第 3 页）：OCR 把三行广告并成
    "拍照批改互动讲解 智能错题本 分钟教材变式练 实际用时："，
    解析器给它单独立了一道题，裁图又顺着彩色横幅裁出 581×77 的页眉废图，
    审计报"第 2 题 切边 上+291px"。
    """
    from app.extractor.rules import is_furniture_text
    lines = ["1.第一题（    ）",
             "拍照批改互动讲解",
             "智能错题本",
             "分钟教材变式练            实际用时：",
             "2.第二题（    ）"]
    got = qs(extract(lines))
    assert len(got) == 2, f"页眉行被当成题目了：{[q.get('text') for q in got]}"
    assert all("拍照批改" not in str(q.get("text")) and "智能错题本" not in str(q.get("text"))
               for q in got), f"页眉文字混进题干：{[q.get('text') for q in got]}"

    assert is_furniture_text("拍照批改互动讲解"), "页眉广告没当家具（会裁成废图）"
    assert is_furniture_text("智能错题本")
    assert is_furniture_text("分钟教材变式练            实际用时：")
    # 反例：题干里恰好出现广告里的词（短句里带"实际用时"）不算页眉
    assert not is_furniture_text("小明做口算题，实际用时比计划少了3分钟。"), \
        "正常题干被当页眉抹掉了"
t("页眉扫码批改广告不是题目（含反例）", page_header_ads_are_not_questions)


def _ann(text, page=1, x=20.0, y=0.0, h=18.0):
    return {"text": text, "page": page, "x": x, "y": y, "w": 260.0, "h": h}


def decimal_like_question_numbers():
    """题号后紧跟数字（2.78÷3、4.3600000000=）是真题的题首，不是小数续行。

    新下载的 P2/P4/P5 实测：旧版为防 "3.5 元" 把这类行一律当续行，
    2/3/4/5/7 题整段并进上一题，用户看到的就是"题目不全"。
    """
    lines = ["一、填空",
             "1.34×7这个算式可以读作（ ），也可以读作（  ）",
             "2.78÷3这个算式可以读作（  ），也可以读作（  ）",
             "3.2米=（  ）分米   70厘米=（  ）分米",
             "4.6000米=（  ）千米  4000千克=（  ）吨",
             "5.60分米=（  ）厘米  3吨=（  ）千克"]
    ann = [_ann(t, y=20 + i * 30) for i, t in enumerate(lines)]
    got = qs(extract(lines, ann=ann))
    assert [q["q_no"] for q in got] == [1, 2, 3, 4, 5], \
        f"题号没分开：{[(q['q_no'], q['text'][:12]) for q in got]}"
    assert got[1]["text"].startswith("78÷3"), f"第2题题干丢了：{got[1]['text']}"
    assert got[3]["text"].startswith("6000米"), f"第4题题干丢了：{got[3]['text']}"

    # 反例①：缩进 40pt 的 "2.78÷3" 是续行（起笔没对齐），不许另起一题
    lines2 = ["一、填空", "1.34×7这个算式可以读作（  ）", "   2.78÷3这个算式可以读作（  ）"]
    ann2 = [_ann(lines2[0], y=20), _ann(lines2[1], y=50),
            _ann(lines2[2], x=60.0, y=80)]
    got2 = qs(extract(lines2, ann=ann2))
    assert [q["q_no"] for q in got2] == [1], f"缩进续行被当成新题：{got2}"
    # 反例②：和上一行同一 y（OCR 把一行切成两段）也不许另起一题
    ann3 = [_ann(lines2[0], y=20), _ann(lines2[1], y=50),
            _ann(lines2[2], y=50)]
    got3 = qs(extract(lines2, ann=ann3))
    assert [q["q_no"] for q in got3] == [1], f"同一行的碎片被当成新题：{got3}"
    # 反例③：对齐、换行，但数字不接上一题（1 后面直接 3.5）——是真小数续行
    lines4 = ["一、填空", "1.34×7的积是多少", "3.5 米比 4 米短多少"]
    ann4 = [_ann(t, y=20 + i * 30) for i, t in enumerate(lines4)]
    got4 = qs(extract(lines4, ann=ann4))
    assert [q["q_no"] for q in got4] == [1], f"不接号的 3.5 被当成新题：{got4}"
t("题号后紧跟数字要认成新题（含三个反例）", decimal_like_question_numbers)


def instruction_lines_are_not_questions():
    """卷首/密封框的考试说明不是题（新下载 P1/P5 实测：第一页凭空多 4 道题）。"""
    lines = ["1.本练习卷共8页，共五道大题。作答时长90分钟。",
             "2.在练习卷和答题纸上准确填写学校、班级和姓名。",
             "3.请仔细读题，按题目要求在答题纸相应位置作答。注意书写清晰。",
             "4.作答结来后按照学校具体要求完成收交。",
             "一、选择题（共2题）",
             "1.3+4=（  ）",
             "2.5+6=（  ）"]
    got = qs(extract(lines))
    assert len(got) == 2, f"考试说明被当成题：{[(q['q_no'], q['text'][:16]) for q in got]}"
    assert [q["q_no"] for q in got] == [1, 2]
    # 反例：长题干里引用"请仔细读题…答题"不能整行丢掉（限长 90 字）
    long_q = ("请仔细读题，把下面这段话中关于答题顺序的要求用横线画出来，再说说为什么先做容易的题"
              "再做难的题，这样安排对考试有什么好处；如果你在考场上遇到一时想不起来的题目，"
              "应该先跳过去做后面的题，还是停下来反复想，请把你的做法和理由写在本子上，"
              "并在小组里说一说，看看谁的办法更合理。")
    assert len("".join(long_q.split())) > 90, "反例本身不够长，测不到限长判据"
    got2 = qs(extract([long_q]))
    assert got2 and got2[0]["text"] == long_q, "长题干被考试说明判据吞掉了"
t("考试说明不进题目（含长题干反例）", instruction_lines_are_not_questions)


def subsection_heading_is_not_a_question():
    """「4、估算」这种只有运算名的编号行是小节标题，不是第 4 题（新下载 P2 实测重号）。"""
    lines = ["一、口算", "1. 25+17=", "2. 40-8=", "3. 6×7=", "4、估算", "5. 39×21≈"]
    r = extract(lines)
    got = qs(r)
    assert [q["q_no"] for q in got] == [1, 2, 3, 5], \
        f"小节标题被当成题：{[(q['q_no'], q['text']) for q in got]}"
    assert len(r["stage_markers"]) == 1, f"小节标题没进 stage_markers：{r['stage_markers']}"
t("只有运算名的编号行是小节标题（含反例）", subsection_heading_is_not_a_question)


def calc_section_circled_items_are_subitems():
    """计算大题的 ①②③④ 是并排小题，不是 A/B/C/D 选项。

    新下载 P2 第 3 页、P3 第 2 页实测：递等式/解方程的 4 个算式被拆成选项，
    题目变成"选一个算式"，做题页只给单选框，孩子没有列竖式/写步骤的地方。
    """
    lines = ["三、计算（30分）",
             "2.递等式计算（12分）",
             "①0.125×9×8×0.3            ②4.62÷5÷0.66",
             "③5.78×2.3+5.78+5.78×6.7    ④7.75×[20÷（3.24-3.04）]",
             "3.解方程（12分）",
             "①6X+15+7=141            ②15X+6X=168"]
    r = extract(lines)
    got = qs(r)
    q2 = [q for q in got if "递等式" in q["text"]][0]
    assert not q2["options"], f"计算小题被拆成选项：{q2['options']}"
    assert q2["qtype"] != "choice", f"计算题被当成选择题：{q2['qtype']}"
    for expr in ("0.125×9×8×0.3", "4.62÷5÷0.66", "7.75×[20÷"):
        assert expr in q2["text"], f"小题丢了：{expr} / {q2['text']}"
    q3 = [q for q in got if "解方程" in q["text"]][0]
    assert not q3["options"] and "15X+6X=168" in q3["text"], q3

    # 反例：真正的选择题里圆圈数字仍是选项
    lines2 = ["三、选择题", "1.下面哪个方向是东？（  ）",
              "①东 ②南 ③西 ④北"]
    q = qs(extract(lines2))[0]
    assert q["qtype"] == "choice" and len(q["options"]) == 4, q

    # OCR 把小标题和第一行算式并成一行：也不许拆成选项
    lines3 = ["三、计算（30分）",
              "2.递等式计算：①0.125×9×8×0.3  ②4.62÷5÷0.66"]
    q3 = qs(extract(lines3))[0]
    assert not q3["options"] and "0.125" in q3["text"], q3

    # 题干里带"选择"字样（"选择合适的方法"）也不许把计算小题拆成选项
    lines4 = ["三、计算（30分）",
              "2.用简便方法计算，选择最合适的一种：①0.125×9×8×0.3  ②4.62÷5÷0.66"]
    q4 = qs(extract(lines4))[0]
    assert not q4["options"] and "0.125" in q4["text"], q4
t("计算大题的圆圈小题不拆成选项（含选择题反例）", calc_section_circled_items_are_subitems)


def fill_section_circled_items_are_subitems():
    """填空题里的 ①②③ 是并列小问，不是四选一（P2 Q8 独立复测）。

    原卷："在□里填上>、<或=" 下面六行比较项（⑤⑥ 是图）。旧版把 ①–④ 拆成
    四个选项，整题变成单选，六个空只能选一个，⑤⑥ 连控件都没有。
    """
    lines = ["一、填空（20分，每空1分）",
             "8. 在□里填上“>”、“<”或“=”。",
             "①23×0□23+0            ②32×2□76",
             "③67厘米□7分米            ④989克□1千克"]
    q = qs(extract(lines))[0]
    assert not q["options"], f"比较项被拆成选项：{q['options']}"
    assert q["qtype"] != "choice", f"填空题被当成选择题：{q['qtype']}"
    for item in ("①23×0", "②32×2", "③67厘米", "④989克"):
        assert item in q["text"], f"比较项丢了：{item} / {q['text']}"
    # 六个空都得有地方写：退成问答（textarea）而不是单选
    assert q["qtype"] in ("fill", "essay"), q["qtype"]

    # 反例①：真正的大题「选择」里圆圈仍是选项
    lines2 = ["三、选择题", "1.下面哪个方向是东？（  ）",
              "①东 ②南 ③西 ④北"]
    q2 = qs(extract(lines2))[0]
    assert q2["qtype"] == "choice" and len(q2["options"]) == 4, q2

    # 反例②：填空大题里题干要求"选出"时，圆圈仍是选项（按题干判，不按大题名一刀切）
    lines3 = ["一、填空（20分）",
              "3. 选出与其它三个不同类的图形（  ）",
              "①三角形 ②正方形 ③圆形 ④长方形"]
    q3 = qs(extract(lines3))[0]
    assert q3["qtype"] == "choice" and len(q3["options"]) == 4, q3
t("填空题的圆圈比较项不拆成四选一（含反例）", fill_section_circled_items_are_subitems)


def bare_section_heading_without_numeral():
    """OCR 把「一、」整段读丢时，标题仍是标题，计分表头也不能变成题。

    新下载 P3 第 1 页实测：「一、填空（20分）」只剩「填空（20分）」，
    产品把它当成一道 1 分的题，后面 13 道填空全部挂在空 section 下；
    P6 第 1 页更只剩「、直接写出得数。（12%）」。
    """
    lines = ["人教版小学五年级数学期末考试卷",
             "填空（20分）",
             "二            三            四            五            总分",
             "1. 2.15小时=（            ）分",
             "2. 3.6吨=（            ）千克",
             "二、判断（10分）",
             "1. 两个面积相等的三角形一定可以拼成一个平行四边形。"]
    r = extract(lines)
    assert r["sections"][0]["name"] == "一、填空（20分）", \
        f"丢号标题没认成大题：{r['sections'][0]['name']!r}"
    assert [q["q_no"] for q in r["sections"][0]["questions"]] == [1, 2], \
        f"题目挂错了段：{[(q['q_no'], q['text']) for q in r['sections'][0]['questions']]}"
    assert r["sections"][1]["name"].startswith("二、"), r["sections"][1]["name"]
    assert all("总分" not in (q.get("text") or "") for q in qs(r)), \
        f"计分表头被当成题：{[q['text'] for q in qs(r)]}"

    # 只剩一个顿号：「、直接写出得数。（12%）」也要认成标题，不是无号题
    r2 = extract(["、直接写出得数。（12%）",
                  "27+9=        20+35=        60×7=        450-20=        76-40=",
                  "二、填空（10分）",
                  "1. 一年有（  ）个月。"])
    assert r2["sections"][0]["name"] == "一、直接写出得数。（12%）", \
        f"顿号开头的标题没认出来：{r2['sections'][0]['name']!r}"
    assert r2["sections"][0]["questions"][0]["text"].startswith("27+9="), \
        f"口算行挂错段：{[q['text'] for q in qs(r2)]}"
    assert r2["sections"][1]["questions"][0]["q_no"] == 1, r2["sections"][1]

    # 反例：正文里的普通「计算下面各题」没有分值括号，不许当大题
    r3 = extract(["一、计算", "1. 计算下面各题。", "25+17=        40-8="])
    assert len(r3["sections"]) == 1 and r3["sections"][0]["questions"], r3["sections"]
t("丢号的「填空（20分）」仍是大题（计分表头不当题）", bare_section_heading_without_numeral)


def numbered_stage_with_content_becomes_a_question():
    """「4、估算」后面跟着没编号的算式 → 它是一道题，内容不许并进上一题。

    新下载 P2 第 3 页实测：估算的 4 个算式被并进第 3 题"用递等式计算"，
    第 4 题凭空消失；若紧跟着第 5 题，则它只是小节标题（见上一条）。
    """
    lines = ["四、计算(30分)",
             "3．用递等式计算(12分)",
             "①34÷6            ②659×8",
             "4、估算",
             "901×8≈        98÷5≈        989+605≈"]
    ann = [_ann(t, y=30 + i * 30) for i, t in enumerate(lines)]
    got = qs(extract(lines, ann=ann))
    assert [q["q_no"] for q in got] == [3, 4], \
        f"估算题号丢了/并进上一题：{[(q['q_no'], q['text']) for q in got]}"
    assert "901×8" in got[1]["text"] and "估算" in got[1]["text"], got[1]
    assert "901×8" not in got[0]["text"], f"估算算式仍并在上题：{got[0]['text']}"
t("带内容的编号小节标题是一道题（估算题不许丢）", numbered_stage_with_content_becomes_a_question)


def same_page_wrong_number_is_repaired():
    """同一页题号倒退/跳号多半是 OCR 读错数字，不是新的编号序列。"""
    # 第 8 题被读成 5.（P4 实测）：同页倒退 → 接上一题改号，且不开"·续"
    lines = ["一、填空"] + [f"{i}.第{i}题" for i in range(1, 8)] + ["5.第八题", "9.第九题"]
    ann = [_ann(t, y=20 + i * 30) for i, t in enumerate(lines)]
    r = extract(lines, ann=ann)
    got = qs(r)
    assert [q["q_no"] for q in got] == list(range(1, 10)), \
        f"同页错号没修：{[q['q_no'] for q in got]}"
    assert len(r["sections"]) == 1, f"同页错号多开了小节：{[s['name'] for s in r['sections']]}"

    # 第 3 题被读成 9.（P4 实测）：后一题又回到 4 → 说明 9 是读错，不是重编号
    lines2 = ["一、填空", "1.第一题", "2.第二题", "9.第三题", "4.第四题"]
    ann2 = [_ann(t, y=20 + i * 30) for i, t in enumerate(lines2)]
    got2 = qs(extract(lines2, ann=ann2))
    assert [q["q_no"] for q in got2] == [1, 2, 3, 4], \
        f"读大的题号没修：{[q['q_no'] for q in got2]}"

    # 反例：跨页倒退是教辅"第N天"的新编号序列，仍要开"·续"
    lines3 = ["一、填空", "1.第一题", "2.第二题", "1.新一天第一题"]
    ann3 = [_ann(lines3[0], page=1, y=20), _ann(lines3[1], page=1, y=50),
            _ann(lines3[2], page=1, y=80), _ann(lines3[3], page=2, y=20)]
    r3 = extract(lines3, ann=ann3)
    assert len(r3["sections"]) == 2, f"跨页倒退没分段：{r3['sections']}"
    assert r3["sections"][1]["name"].endswith("·续"), r3["sections"][1]["name"]
t("同页 OCR 错号要修、跨页倒退才分段（含反例）", same_page_wrong_number_is_repaired)


def inline_merged_column_questions_split():
    """两栏并排被 OCR 拼成一行："3. 3.6吨…  4.0.8公顷…" 要拆成两道题。"""
    lines = ["一、填空",
             "1. 2.15小时=（  ）分",
             "2.（  ）米=18厘米",
             "3. 3.6吨=（  ）千克    4.0.8公顷=（  ）平方米",
             "5.4.15×0.53的积有（  ）位小数"]
    ann = [_ann(t, y=20 + i * 30) for i, t in enumerate(lines)]
    got = qs(extract(lines, ann=ann))
    assert [q["q_no"] for q in got] == [1, 2, 3, 4, 5], \
        f"并排题没拆开：{[(q['q_no'], q['text'][:18]) for q in got]}"
    assert "0.8公顷" in got[3]["text"], f"第4题正文不对：{got[3]['text']}"
    assert "0.8公顷" not in got[2]["text"], f"第4题还并第3题里：{got[2]['text']}"

    # 反例：正文里只隔一个空格的 "4.5 与…" 是小数，不是题号 4
    lines2 = ["一、填空", "1.第一题", "2.第二题", "3. 3.6吨=（  ）千克 4.5 与 8.1 比较"]
    ann2 = [_ann(t, y=20 + i * 30) for i, t in enumerate(lines2)]
    got2 = qs(extract(lines2, ann=ann2))
    assert [q["q_no"] for q in got2] == [1, 2, 3], f"小数被拆成新题：{got2}"
    assert "4.5" in got2[2]["text"], f"题干里的小数丢了：{got2[2]['text']}"
t("OCR 并排题按下一题号拆开（含小数反例）", inline_merged_column_questions_split)


def ann_split_for_side_by_side_questions():
    """一行两题被 OCR 读成一个框：ann 要按文字比例拆成两条、x 分开。

    真实卷 P1 第 17/18 题：文本层早已切成两题，ann 仍是一条整行框 ——
    两题共用同一个 y 锚点，裁图时前者的区间被后者截断，两幅柱状图全丢。
    """
    ann = [{"text": "17. 松树种了多少棵？          18. 科技书有多少本？",
            "page": 1, "col": 0, "x": 20.0, "y": 40.0, "w": 520.0, "h": 12.0}]
    out = split_question_ann(ann)
    assert len(out) == 2, f"应拆成两条：{[a['text'] for a in out]}"
    assert out[0]["text"].startswith("17.") and out[1]["text"].startswith("18."), \
        f"顺序/内容错：{[a['text'] for a in out]}"
    assert out[0]["x"] < out[1]["x"], f"起点没分开：{[a['x'] for a in out]}"
    assert out[0]["w"] > 40 and out[1]["w"] > 40, f"宽度不合理：{[a['w'] for a in out]}"
    assert out[0]["x"] + out[0]["w"] <= 20.0 + 520.0 + 0.1
    assert out[0]["y"] == out[1]["y"] == 40.0 and out[0]["page"] == out[1]["page"]

    # 光杆题号（P3 第 1/2 小题一行："1.        2、"）：序号连续也要拆
    bare = [{"text": "1.                2、", "page": 2, "col": 0,
             "x": 20.0, "y": 437.0, "w": 360.0, "h": 12.0}]
    ob = split_question_ann(bare)
    assert len(ob) == 2 and ob[0]["x"] < ob[1]["x"], f"光杆题号没拆：{ob}"
t("并排两题的 ann 整行框按比例拆开（P1/P3 正例）", ann_split_for_side_by_side_questions)


def ann_split_rejects_non_questions():
    """反例：小数、单个题号、不连续的光杆数字、行首有正文 —— 一律不切。"""
    def one(text):
        return len(split_question_ann([{"text": text, "page": 1, "col": 0,
                                        "x": 20.0, "y": 20.0, "w": 300.0, "h": 12.0}]))
    assert one("3.5 与 4.2 哪个大？") == 1, "小数被当两个题号"
    assert one("17. 松树种了多少棵？") == 1, "单题被拆"
    assert one("1.          3、") == 1, "不连续的光杆数字被拆"
    assert one("答：17. 第一问          18. 第二问") == 1, "行首有正文也拆"
    assert one("1.      2. 有正文") == 1, "一光杆一有正文被拆"
t("ann 拆行不误伤小数/单题/不连续编号（反例）", ann_split_rejects_non_questions)


def side_by_side_spans_are_not_cut_by_each_other():
    """并排两题的区间下界要拿【本题自己的起点行】比，不能拿向上扩过的 y0。

    旧判据 `o.y > y0 + 1` 在 y0 被抬到大题题头后，把同一行的另一道题也当成了
    "下面的锚点"，两题区间都只剩题头+题干 —— P1 第 17/18 题各丢一张图。
    """
    lines = ["四、看图列式计算（共2道小题）",
             "17. 松树种了多少棵？          18. 科技书有多少本？",
             "五、解决问题（20分）"]
    r = extract(lines, "并排题")
    q = qs(r)
    assert [x["q_no"] for x in q] == [17, 18], \
        f"文本层没切成两题：{[(x['q_no'], x['text']) for x in q]}"
    ann = split_question_ann([
        {"text": lines[0], "page": 1, "col": 0, "x": 20.0, "y": 10.0,
         "w": 300.0, "h": 12.0},
        {"text": lines[1], "page": 1, "col": 0, "x": 20.0, "y": 40.0,
         "w": 520.0, "h": 12.0},
        {"text": lines[2], "page": 1, "col": 0, "x": 20.0, "y": 200.0,
         "w": 300.0, "h": 12.0},
    ])
    assert len(ann) == 4, f"ann 没拆：{[a['text'] for a in ann]}"
    spans = spans_for_questions(
        ann, q, pages=[{"page": 1, "width_pt": 600.0, "height_pt": 800.0, "bands": []}],
        sections=r["sections"])
    by = {int(s["q_index"]): s for s in spans}
    assert set(by) == {0, 1}, f"两题都要有区间：{sorted(by)}"
    assert by[0]["y1"] >= 190 and by[1]["y1"] >= 190, \
        f"区间被并排的另一题截断：{[(k, v['y0'], v['y1']) for k, v in by.items()]}"
    assert by[0]["y0"] <= 10.5 and by[1]["y0"] <= 10.5, \
        f"区间没向上够到大题题头：{[(k, v['y0']) for k, v in by.items()]}"
t("并排两题的区间互不截断、都够到大题题头", side_by_side_spans_are_not_cut_by_each_other)


def docx_scanned_pages_become_scan_pages():
    """扫描页贴进 Word：每张整页图要当扫描页，不许只当 page_count=1 的配图。

    新下载 P1/P2/P3 全是这种 DOCX：旧版识别 0 行 0 题，预览却能出 4–6 页。
    """
    import tempfile
    from docx import Document
    from docx.shared import Inches
    from PIL import Image
    from app.parsers.docx import parse_docx
    from app.extractor.lines import _from_parsed_docx

    with tempfile.TemporaryDirectory() as td:
        img1 = os.path.join(td, "p1.png")
        img2 = os.path.join(td, "p2.png")
        Image.new("RGB", (600, 800), "white").save(img1)
        Image.new("RGB", (580, 780), "white").save(img2)
        doc = Document()
        doc.add_picture(img1, width=Inches(4))
        doc.add_picture(img2, width=Inches(4))
        path = os.path.join(td, "scan.docx")
        doc.save(path)
        r = parse_docx(path, media_dir=os.path.join(td, "media"))
        assert r["page_count"] == 2, f"页数不对：{r['page_count']}"
        assert r["scan_pages"] == 2, f"扫描页数不对：{r['scan_pages']}"
        assert [p["is_scan"] for p in r["pages"]] == [True, True]
        assert all(os.path.isfile(p["rendered_image"]) for p in r["pages"]), \
            "扫描页没有落出整页图"
        # 文档流顺序：第一页应是 img1（旧版按包内顺序会颠倒/错页）
        with Image.open(r["pages"][0]["rendered_image"]) as im0:
            assert im0.size == (600, 800), f"页序错了：{im0.size}"

        # 文字 DOCX 也必须返回 ann；旧版只给 lines，任务链路取 ann 时整份文字蒸发
        text_parsed = {"pages": [{"page_no": 1, "is_scan": False,
                                  "blocks": [{"kind": "text", "text": "1.第一题\n2.第二题"}]}]}
        t2 = _from_parsed_docx({"parsed": text_parsed})
        assert t2["lines"] == ["1.第一题", "2.第二题"], t2["lines"]
        assert len(t2["ann"]) == 2 and t2["ann"][1]["y"] > t2["ann"][0]["y"], t2["ann"]
t("Word 里的扫描页按页 OCR、文字页带坐标（含页序）", docx_scanned_pages_become_scan_pages)


def wrapped_circled_option_keeps_its_tail():
    """圆圈说法折行要接回说法，不能并进题干或被切断。

    新下载 P5 Q7 实测：③ 被 OCR 拦腰断成两行，第二行 "命中率一定也是40%"
    没有编号，旧版并回题干 → ③ 只剩半句（P1 级）。现在 ①②③ 本身是
    "说法"，正确落点是题干（A-D 才是可勾选项），折行必须一起带过去。
    """
    lines = [
        "一、选择题（每小题只有1个正确选项）",
        "7. 淘气正在参加学校组织的一场套圈游戏，目前他已经套圈20次，这20次他的套圈",
        "命中率为40%。以下三种说法中，正确的（）。",
        "①在目前这20次套圈中，淘气套中了8次：",
        "②在目前这20次套圈中，淘气没套中的次数占60%：",
        "③在这场游戏中，如果淘气再接着套圈20次，那么接下来这20次他的套圈",
        "命中率一定也是40%",
        "A.只有②         B.只有①②         C.只有①③         D.只有②③",
    ]
    q = qs(extract(lines, "测试卷"))
    assert len(q) == 1, f"应只有 1 题，实际 {len(q)}"
    assert "命中率一定也是40%" in q[0]["text"], f"③ 的折行丢了：{q[0]['text']}"
    assert "①在目前这20次套圈中，淘气套中了8次" in q[0]["text"], q[0]["text"]
    assert [o[0] for o in q[0]["options"]] == ["A", "B", "C", "D"], \
        f"说法被印成可勾选项：{q[0]['options']}"

    # 反例一：上一句已经有句末标点 → 下一行不能塞进选项
    lines2 = [
        "一、选择题",
        "1. 下面说法正确的是（）。",
        "①圆的直径是半径的2倍。",
        "请选择正确的选项。",
        "A.①         B.①②",
    ]
    q2 = qs(extract(lines2, "测试卷"))
    assert "请选择正确的选项" in q2[0]["text"], f"折行判断误伤题干：{q2[0]['text']}"
    assert all("请选择" not in o for o in q2[0]["options"]), q2[0]["options"]

    # 反例二：计算大题的圆圈小题，后面的算式行不许被吞进上一条
    lines3 = [
        "三、计算题",
        "1. 直接写出得数。",
        "①2.8:0.7              ②3.54:12",
        "2. 脱式计算",
        "①1.25×8                ②4.8÷0.6",
    ]
    q3 = qs(extract(lines3, "测试卷"))
    assert all(not x["options"] for x in q3), \
        f"计算小题被当成圆圈选项：{[(x['text'], x['options']) for x in q3]}"
    all_text = " ".join(x["text"] or "" for x in q3)
    assert "①2.8:0.7" in all_text and "①1.25×8" in all_text, \
        f"算式被吞掉：{all_text}"
t("圆圈选项折行接回选项（含两个反例）", wrapped_circled_option_keeps_its_tail)


def bare_circled_marker_is_not_an_option():
    """A-D 选项旁孤零零的圆圈序号是图注残渣，不是选项。

    新下载 P5 Q3 实测：影子图里单独一行的 "③" 被当成选项，卷面多出一个
    "☐③"；P1 Q7 更是一口气多了 ②③④ 三个。真实选项是 A-D 时必须清掉。
    """
    lines = [
        "一、选择题",
        "1. 下面说法正确的是（  ）",
        "③",
        "A.只有①         B.只有②         C.只有③         D.只有④",
    ]
    q = qs(extract(lines, "测试卷"))
    assert [o[0] for o in q[0]["options"]] == ["A", "B", "C", "D"], \
        f"图注残渣混进选项：{q[0]['options']}"

    # 反例：整题选项就是圆圈条目（没有 A-D）→ 一条都不许删
    lines2 = [
        "一、选择题",
        "1. 下面说法正确的是（  ）",
        "①圆的直径是半径的2倍",
        "②半圆的周长是圆周长的一半",
    ]
    q2 = qs(extract(lines2, "测试卷"))
    stripped = [o for o in q2[0]["options"] if o.strip() not in ("①", "②")]
    assert len(stripped) >= 1, f"圆圈选项被误删：{q2[0]['options']}"

    # 反例二：三条圆圈说法 + A-D 选答案（P5 Q7）→ 说法挪回题干，选项只留 A-D
    lines3 = [
        "一、选择题",
        "1. 以下三种说法中，正确的（  ）",
        "①说法一",
        "②说法二",
        "③说法三的结论",
        "A.只有②         B.只有①②         C.只有①③         D.只有②③",
    ]
    q3 = qs(extract(lines3, "测试卷"))
    assert [o[0] for o in q3[0]["options"]] == ["A", "B", "C", "D"], \
        f"三条说法还当成选项（会被印成三个 ☐）：{q3[0]['options']}"
    assert all(x in q3[0]["text"] for x in ("说法一", "说法二", "说法三的结论")), \
        f"说法被删掉了：{q3[0]['text']!r}"

    # 正例二：只有一条圆圈图注 + A-D（P5 Q10）→ 图注要清掉
    lines4 = [
        "一、选择题",
        "1. 下面四个情境中，两个量之比可以用2：3表示的（  ）",
        "②付款20元购买30米线绳",
        "A.只有①④         B.只有②④         C.只有①②④         D.只有①②③",
    ]
    q4 = qs(extract(lines4, "测试卷"))
    assert [o[0] for o in q4[0]["options"]] == ["A", "B", "C", "D"], \
        f"单条圆圈图注混进选项：{q4[0]['options']}"
t("A-D 旁单独的圆圈序号不是选项（含四个反例）", bare_circled_marker_is_not_an_option)


print(f"\n{ok_n} passed, {bad_n} failed")
sys.exit(1 if bad_n else 0)
